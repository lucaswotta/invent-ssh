"""Regression tests for export quality, metadata preservation, and safe rendering."""
import csv
import importlib.util
import tempfile
import unittest

from reporting import save_report


SAMPLE = {
    "ip": "192.0.2.10",
    "metadata": {
        "HOSTNAME": "lab-node-01",
        "ASSET_TAG": "ASSET-010",
        "LOCATION": "Test Lab",
        "OWNER": "Infrastructure",
        "NOTES": "<script>alert(1)</script>",
        "FORMULA_TEST": "=HYPERLINK(\"https://example.invalid\",\"click\")",
    },
    "status": "ONLINE",
    "placa_mae": "Example Board",
    "processador": "Example CPU",
    "cores_threads": "8/16",
    "ram": "32GB DDR4",
    "disk_type": "NVMe",
    "disk_size": "1TB",
    "distro": "Ubuntu 24.04 LTS",
    "kernel": "6.8",
    "erro": "",
    "dta_atualizacao": "2026-10-19 14:00:00",
    "hardware_details": {
        "hostname": "detected-lab-01",
        "manufacturer": "Example Systems",
        "model": "Model A",
        "architecture": "x86_64",
        "storage_devices": "nvme0n1 · NVMe · 512 GB · Example Drive",
        "gpu": "Example Integrated GPU",
        "network_adapters": "enp1s0 (up)",
        "uptime": "5d 3h 12m",
        "system_serial": "Not collected (opt-in)",
        "coverage_percent": "92%",
        "collection_notes": "GPU model not available",
    },
}


class ReportingTests(unittest.TestCase):
    def test_html_is_self_contained_searchable_and_escapes_untrusted_values(self):
        with tempfile.TemporaryDirectory() as directory:
            path = save_report([SAMPLE], "HTML", directory)
            document = path.read_text(encoding="utf-8")

        self.assertIn("Hardware inventory", document)
        self.assertIn('id="search"', document)
        self.assertIn("lab-node-01", document)
        self.assertIn("ASSET-010", document)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", document)
        self.assertIn("&lt;script&gt;", document)
        self.assertNotIn("<script>alert(1)</script>", document)
        self.assertIn("data-status=\"ONLINE\"", document)

    def test_csv_preserves_custom_columns_and_neutralizes_formula_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = save_report([SAMPLE], "CSV", directory)
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                parsed = list(csv.DictReader(handle, delimiter=";"))

        self.assertEqual(len(parsed), 1)
        row = parsed[0]
        self.assertEqual(row["Host / IP"], "192.0.2.10")
        self.assertEqual(row["Hostname"], "lab-node-01")
        self.assertEqual(row["Asset tag"], "ASSET-010")
        self.assertEqual(row["Location"], "Test Lab")
        self.assertEqual(row["Status"], "ONLINE")
        self.assertEqual(row["Detected hostname"], "detected-lab-01")
        self.assertEqual(row["Manufacturer"], "Example Systems")
        self.assertIn("nvme0n1", row["Storage devices"])
        self.assertEqual(row["Data coverage"], "92%")
        self.assertTrue(row["Formula Test"].startswith("'="))
        self.assertIn("<script>alert(1)</script>", row["Notes"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is not installed")
    def test_xlsx_has_overview_table_filters_freeze_panes_and_safe_text(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = save_report([SAMPLE], "XLSX", directory)
            workbook = load_workbook(path, data_only=False)
            self.addCleanup(workbook.close)

            self.assertEqual(workbook.sheetnames, ["Overview", "Inventory"])
            inventory = workbook["Inventory"]
            self.assertEqual(inventory.freeze_panes, "A5")
            self.assertEqual(inventory.auto_filter.ref, f"A4:{inventory.cell(4, inventory.max_column).column_letter}{inventory.max_row}")
            self.assertGreaterEqual(len(workbook["Overview"]._images), 1)
            self.assertGreaterEqual(len(inventory._images), 1)
            headers = [cell.value for cell in inventory[4]]
            notes_column = headers.index("Notes") + 1
            notes_value = inventory.cell(5, notes_column).value
            self.assertIn("<script>alert(1)</script>", notes_value)
            formula_column = headers.index("Formula Test") + 1
            formula_cell = inventory.cell(5, formula_column)
            self.assertNotEqual(formula_cell.data_type, "f")
            self.assertTrue(str(formula_cell.value).startswith("'="))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is not installed")
    def test_xlsx_report_respects_portuguese_locale(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = save_report([SAMPLE], "XLSX", directory, language="pt-br")
            workbook = load_workbook(path, data_only=False)
            self.addCleanup(workbook.close)

            self.assertEqual(workbook.sheetnames, ["Resumo", "Inventário"])
            inventory = workbook["Inventário"]
            headers = [cell.value for cell in inventory[4]]
            self.assertIn("Nome do host", headers)
            self.assertIn("Localização", headers)
            self.assertIn("Responsável", headers)
            self.assertIn("Placa-mãe", headers)

    def test_html_report_respects_portuguese_locale(self):
        with tempfile.TemporaryDirectory() as directory:
            path = save_report([SAMPLE], "HTML", directory, language="pt-br")
            document = path.read_text(encoding="utf-8")

        self.assertIn('<html lang="pt-br">', document)
        self.assertIn("Inventário de hardware", document)
        self.assertIn("Total de máquinas", document)
        self.assertIn("Localização", document)
        self.assertIn("máquinas inventariadas", document)
        self.assertIn("Todos os status", document)
        self.assertIn("Não coletado (opcional)", document)
        self.assertIn("Gerado localmente pelo Invent-SSH", document)
        self.assertIn("Acessível</span>", document)

    def test_metadata_names_colliding_with_hardware_columns_are_preserved(self):
        from reporting import _prepare

        record = dict(SAMPLE)
        record["metadata"] = {
            "STATUS": "asset-side-status",
            "ARCHITECTURE": "asset-side-architecture",
            "HOSTNAME": "registered-host",
        }
        columns, rows = _prepare([record], "en")
        row = rows[0]

        self.assertIn("Asset field · Status", columns)
        self.assertEqual(row["Asset field · Status"], "asset-side-status")
        self.assertIn("Asset field · Architecture", columns)
        self.assertEqual(row["Asset field · Architecture"], "asset-side-architecture")
        self.assertEqual(row["Hostname"], "registered-host")
        self.assertEqual(row["Detected hostname"], "detected-lab-01")
        self.assertEqual(row["Status"], "ONLINE")

    def test_unsupported_format_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                save_report([SAMPLE], "PDF", directory)

    def test_empty_report_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                save_report([], "HTML", directory)


if __name__ == "__main__":
    unittest.main()
