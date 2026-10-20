"""
app.py: A Interface Gráfica Principal (GUI) do invent-ssh.

Este módulo contém a classe `App`, que orquestra a interface do usuário
construída com CustomTkinter. Ele gerencia as diferentes abas de operação
(lista de ativos), coleta as configurações do usuário e inicia o motor
de inventário (`InventoryEngine`).
"""
# --- Módulos Padrão e de Terceiros ---
import sys
import subprocess
import threading
import queue
import json
import os
import glob
import platform
import traceback
import re
import webbrowser
import logging
from datetime import datetime
try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, font as tkfont
except ImportError:  # Minimal Linux/Python images may omit Tcl/Tk.
    tk = None
    filedialog = None
    messagebox = None
    tkfont = None

from typing import Optional, Dict, Any

# i18n depende apenas da stdlib, então pode ser importado incondicionalmente
# mesmo antes da checagem de dependências de terceiros.
from i18n import t, set_language, get_language, DEFAULT_LANGUAGE, LANGUAGE_DISPLAY_NAMES
from paths import APP_DISPLAY_NAME, get_config_path, get_logs_dir, save_user_config

# --- Importações condicionais para dependências ---
# Esta estrutura garante que a verificação de dependências ocorra
# antes de tentar importar os módulos necessários para a GUI.
try:
    import customtkinter as ctk
    from openpyxl import Workbook
    from PIL import Image as PILImage
    from core import InventoryEngine
except ImportError:
    # As importações serão tratadas pelo verificador de dependências.
    # Se falhar, a aplicação não continuará.
    ctk = None
    Workbook = None
    PILImage = None
    InventoryEngine = None

# --- Constantes da Aplicação ---
APP_VERSION = "0.4.0"  # Public rebrand and cross-distro collector
APP_BRAND = APP_DISPLAY_NAME
APP_NAME = "invent-ssh"  # Stable filesystem/log prefix.
CONFIG_FILE = str(get_config_path())
ICON_FILE = "app.ico"

# --- Constantes de Configuração Padrão ---

# --- Visual identity / native-feeling cross-platform design system ---
COLORS = {
    "canvas": "#061327",
    "paper": "#0A1C35",
    "paper_alt": "#102746",
    "ink": "#EDF4FF",
    "muted": "#9DB2D0",
    "line": "#183B63",
    "pine": "#07172D",
    "forest": "#163B65",
    "sage": "#6BA7E8",
    "sage_soft": "#102A4A",
    "clay": "#FFC629",
    "clay_hover": "#E6A800",
    "sand": "#FFE28A",
    "danger": "#F87171",
    "danger_soft": "#341E2A",
}

THEME = {
    "font_family": "Avenir Next",
    "font_display_family": "Georgia",
    "font_mono_family": "SF Mono",
    "font_h1": ("Georgia", 27, "normal"),
    "font_h2": ("Avenir Next", 16, "bold"),
    "font_body": ("Avenir Next", 13),
    "font_small": ("Avenir Next", 11),
    "font_label": ("Avenir Next", 12, "bold"),
    "font_mono": ("SF Mono", 11),
    "padding": 20,
    "padding_sm": 12,
    "link_color": COLORS["sage"],
    **COLORS,
}


def configure_typography(root) -> None:
    """Choose installed native UI fonts instead of assuming one bundled typeface."""
    try:
        available = {family.casefold(): family for family in tkfont.families(root)}
    except (tk.TclError, RuntimeError):
        available = {}

    system = platform.system()
    if system == "Darwin":
        body_candidates = ("Avenir Next", "Helvetica Neue", "Arial")
        mono_candidates = ("SF Mono", "Menlo", "Monaco", "Courier New")
    elif system == "Windows":
        body_candidates = ("Segoe UI", "Avenir Next", "Arial")
        mono_candidates = ("Cascadia Mono", "Consolas", "Courier New")
    else:
        body_candidates = ("Noto Sans", "Cantarell", "Liberation Sans", "DejaVu Sans", "TkDefaultFont")
        mono_candidates = ("Noto Sans Mono", "DejaVu Sans Mono", "Liberation Mono", "TkFixedFont")

    def select(candidates, fallback):
        return next((available[name.casefold()] for name in candidates if name.casefold() in available), fallback)

    THEME["font_family"] = select(body_candidates, "TkDefaultFont")
    THEME["font_display_family"] = select(
        ("Georgia", "Baskerville", "Palatino Linotype", "Liberation Serif", "DejaVu Serif", "Times"),
        "TkDefaultFont",
    )
    THEME["font_mono_family"] = select(mono_candidates, "TkFixedFont")
    THEME["font_h1"] = (THEME["font_display_family"], 27, "normal")
    THEME["font_h2"] = (THEME["font_family"], 16, "bold")
    THEME["font_body"] = (THEME["font_family"], 13)
    THEME["font_small"] = (THEME["font_family"], 11)
    THEME["font_label"] = (THEME["font_family"], 12, "bold")
    THEME["font_mono"] = (THEME["font_mono_family"], 11)


def bind_responsive_grid(container, widgets, min_column_width: int = 245, max_columns: int = 3) -> None:
    """Reflow a row of feature cards to the available content width."""
    state = {"columns": 0}

    def arrange(_event=None):
        width = max(1, container.winfo_width())
        columns = max(1, min(max_columns, width // min_column_width))
        if columns == state["columns"]:
            return
        state["columns"] = columns
        for column in range(max_columns):
            container.grid_columnconfigure(
                column, weight=1 if column < columns else 0,
                uniform="responsive_cards" if column < columns else "",
            )
        for index, widget in enumerate(widgets):
            row, column = divmod(index, columns)
            widget.grid(
                row=row, column=column, sticky="nsew",
                padx=(0 if column == 0 else 6, 0 if column == columns - 1 else 6),
                pady=(0, 12),
            )
            container.grid_rowconfigure(row, weight=1)
        for row in range((len(widgets) + columns - 1) // columns, max_columns + 1):
            container.grid_rowconfigure(row, weight=0)

    container.bind("<Configure>", arrange, add="+")
    container.after_idle(arrange)


# --- Funções Auxiliares ---

def resource_path(relative_path: str) -> str:
    """
    Obtém o caminho absoluto para um recurso, garantindo compatibilidade
    com o ambiente de desenvolvimento e o executável do PyInstaller.

    Args:
        relative_path: O caminho relativo do recurso a partir da raiz.

    Returns:
        O caminho absoluto para o recurso.
    """
    try:
        # PyInstaller cria uma pasta temporária e armazena o caminho em _MEIPASS
        base_path = sys._MEIPASS
    except AttributeError:
        # Em ambiente de desenvolvimento, o caminho base é o diretório atual
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)


def load_brand_image(filename: str, size: tuple[int, int]):
    """Load a bundled mark as a HiDPI-aware CTkImage without holding a file handle."""
    if ctk is None or PILImage is None:
        raise RuntimeError("CustomTkinter image dependencies are unavailable.")
    image_path = resource_path(os.path.join("brand", filename))
    with PILImage.open(image_path) as source:
        rgba = source.convert("RGBA").copy()
    return ctk.CTkImage(light_image=rgba, size=size)

def check_and_install_dependencies():
    """
    Verify GUI and runtime requirements before constructing any widgets.
    """
    if tk is None or tkfont is None:
        message = (
            "Invent-SSH needs Tcl/Tk support in this Python installation. "
            "On Debian/Ubuntu, install python3-tk for the Python interpreter used to run the app; "
            "on Fedora, install python3-tkinter. On macOS, use a Python distribution that includes Tcl/Tk."
        )
        print(message, file=sys.stderr)
        return False

    required_packages = {
        'customtkinter': 'customtkinter',
        'openpyxl': 'openpyxl',
        'paramiko': 'paramiko',
        'Pillow': 'PIL',
    }
    try:
        from importlib import util
        missing_packages = [pkg for pkg, mod in required_packages.items() if not util.find_spec(mod)]
    except ImportError:
        # Fallback para ambientes mais antigos
        import pkgutil
        missing_packages = [pkg for pkg, mod in required_packages.items() if not pkgutil.find_loader(mod)]

    if not missing_packages:
        return True

    # Usa o Tkinter nativo para o diálogo, pois o customtkinter pode não estar instalado
    root = tk.Tk()
    root.withdraw()

    msg = t("missing_deps_body", packages=', '.join(missing_packages))

    if messagebox.askyesno(t("missing_deps_title"), msg):
        try:
            # Tenta instalar os pacotes ausentes usando pip
            subprocess.check_call([sys.executable, "-m", "pip", "install", *missing_packages])
            messagebox.showinfo(t("deps_install_success_title"), t("deps_install_success_body"))
        except Exception as e:
            messagebox.showerror(t("deps_install_error_title"), t("deps_install_error_body", error=e, packages=' '.join(missing_packages)))
        return False  # Requer reinicialização
    else:
        messagebox.showwarning(t("deps_skip_warning_title"), t("deps_skip_warning_body"))
        return True # Permite continuar, mas pode falhar

# --- Classes da UI ---

class BaseModal(ctk.CTkToplevel if ctk is not None else object):
    """Classe base para janelas modais, garantindo consistência visual e comportamento."""
    def __init__(self, master, title: str, geometry: str, resizable: bool = True):
        super().__init__(master)
        self.title(title)
        self.geometry(geometry)
        self.resizable(resizable, resizable)
        if resizable:
            self.minsize(560, 380)
        self.transient(master)  # Mantém o modal sempre à frente da janela principal
        self.grab_set()         # Captura o foco, impedindo interação com a janela principal

        icon_path = resource_path(ICON_FILE)
        try:
            if platform.system() == "Windows" and os.path.exists(icon_path):
                self.iconbitmap(icon_path)
            else:
                icon_png = resource_path("brand/invent-ssh-mark-64.png")
                if os.path.exists(icon_png):
                    self._modal_icon = tk.PhotoImage(file=icon_png)
                    self.iconphoto(True, self._modal_icon)
        except tk.TclError:
            logging.debug("Modal icon unavailable on this platform.", exc_info=True)

        self.grid_columnconfigure(0, weight=1)

class WelcomeModal(BaseModal):
    """Modal de boas-vindas com foco na proposta de valor, segurança e usabilidade."""
    def __init__(self, master, config: dict):
        super().__init__(master, t("welcome_title", app_name=APP_BRAND), "920x680", resizable=True)
        self.master = master
        self.config = config

        self.protocol("WM_DELETE_WINDOW", self.on_closing_attempt)
        self.grid_rowconfigure(0, weight=1)

        main_frame = ctk.CTkScrollableFrame(self, fg_color="transparent", corner_radius=0)
        main_frame.grid(row=0, column=0, padx=THEME["padding"], pady=THEME["padding"], sticky="nsew")
        main_frame.grid_columnconfigure(0, weight=1)

        def create_section(parent, title, content):
            ctk.CTkLabel(parent, text=title, font=THEME["font_h2"], anchor="w").pack(fill="x", pady=(15, 5), padx=10)
            ctk.CTkLabel(parent, text=content, font=THEME["font_body"], anchor="w", justify="left", wraplength=680).pack(fill="x", padx=10, pady=(0, 10))

        create_section(main_frame, t("welcome_section1_title"), t("welcome_section1_body"))
        create_section(main_frame, t("welcome_section2_title"), t("welcome_section2_body"))
        create_section(main_frame, t("welcome_section3_title"), t("welcome_section3_body"))

        footer_frame = ctk.CTkFrame(self, fg_color="transparent")
        footer_frame.grid(row=1, column=0, padx=THEME["padding"], pady=(0, THEME["padding"]), sticky="ew")
        footer_frame.grid_columnconfigure(0, weight=1)

        self.show_again_var = ctk.BooleanVar(value=True)
        ctk.CTkCheckBox(footer_frame, text=t("welcome_checkbox_dont_show"), variable=self.show_again_var, onvalue=False, offvalue=True, font=THEME["font_body"]).grid(row=0, column=0, padx=5, sticky="w")
        ctk.CTkButton(footer_frame, text=t("welcome_button_start"), command=self.close_modal).grid(row=0, column=1, padx=5, sticky="e")

    def on_closing_attempt(self):
        """Impede o fechamento acidental da janela de boas-vindas."""
        pass

    def close_modal(self):
        """Fecha o modal e salva a preferência do usuário sobre exibi-lo novamente."""
        self.master.config["show_welcome_modal"] = self.show_again_var.get()
        self.destroy()

class Tooltip:
    """Tooltip aprimorado para fornecer ajuda contextual instantânea sobre os widgets."""
    def __init__(self, widget, text: str, delay_ms: int = 800):
        self.widget = widget
        self.text = text
        self.delay = delay_ms
        self.tooltip_window = None
        self.after_id = None
        self.widget.bind("<Enter>", self.schedule_tooltip)
        self.widget.bind("<Leave>", self.hide_tooltip)
        self.widget.bind("<Button-1>", self.hide_tooltip)

    def schedule_tooltip(self, event=None):
        """Agenda a exibição do tooltip após um certo delay."""
        self.hide_tooltip()
        if self.text:
            self.after_id = self.widget.after(self.delay, self._show_tooltip)

    def _show_tooltip(self):
        """Place help below the owning widget without relying on Text-only bbox APIs."""
        if not self.text or not self.widget.winfo_exists():
            return
        try:
            x = self.widget.winfo_rootx() + 4
            y = self.widget.winfo_rooty() + self.widget.winfo_height() + 7
            self.tooltip_window = ctk.CTkToplevel(self.widget)
            self.tooltip_window.wm_overrideredirect(True)
            self.tooltip_window.wm_geometry(f"+{x}+{y}")
            self.tooltip_window.configure(fg_color=THEME["pine"])
            ctk.CTkLabel(
                self.tooltip_window, text=self.text, corner_radius=7,
                fg_color=THEME["pine"], text_color=THEME["paper"],
                font=THEME["font_small"], justify="left",
            ).pack(ipadx=9, ipady=6)
        except tk.TclError:
            self.tooltip_window = None

    def hide_tooltip(self, event=None):
        """Esconde e destrói o tooltip."""
        if self.after_id:
            try:
                self.widget.after_cancel(self.after_id)
            except tk.TclError:
                pass
        if self.tooltip_window:
            try:
                self.tooltip_window.destroy()
            except tk.TclError:
                pass
        self.after_id = None
        self.tooltip_window = None

class LogModal(BaseModal):
    """
    Modal de progresso que exibe logs em tempo real, uma barra de progresso
    e controla a interação do usuário durante a execução do inventário.
    """
    def __init__(self, master):
        super().__init__(master, t("log_modal_title_running"), "960x620", resizable=True)
        self.master = master
        self.log_queue = master.log_queue
        self._modal_finished = False

        self.grid_rowconfigure(1, weight=1)
        self.progress_bar = ctk.CTkProgressBar(self, height=12)
        self.progress_bar.set(0)
        self.progress_bar.grid(row=0, column=0, padx=THEME["padding_sm"], pady=(THEME["padding_sm"], 5), sticky="ew")

        self.log_textbox = ctk.CTkTextbox(self, font=THEME["font_mono"], state="disabled")
        self.log_textbox.grid(row=1, column=0, padx=THEME["padding_sm"], pady=(0, THEME["padding_sm"]), sticky="nsew")

        self.close_button = ctk.CTkButton(self, text=t("log_modal_close_button"), command=self.destroy, state="disabled")
        self.close_button.grid(row=2, column=0, padx=THEME["padding_sm"], pady=(5, THEME["padding_sm"]), sticky="ew")

        self.protocol("WM_DELETE_WINDOW", self.on_closing_attempt)
        self.process_log_queue()

    def process_log_queue(self):
        """Processa mensagens da fila de logs e atualiza a UI periodicamente."""
        try:
            while not self.log_queue.empty():
                level, message, value = self.log_queue.get_nowait()
                if level != "PROGRESS":
                    self.add_log(level, message)

                if level == "PROGRESS" and value is not None:
                    self.progress_bar.set(value / 100)
                elif level == "FINISH":
                    self.finish_process(message, value or "INFO")
                elif level == "OPEN_FILE":
                    self.master.open_file(value)
        finally:
            if not self._modal_finished:
                try:
                    self.after(150, self.process_log_queue)
                except tk.TclError:
                    pass

    def add_log(self, level, message):
        """Adiciona uma nova linha de log ao textbox."""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_textbox.configure(state="normal")
        self.log_textbox.insert("end", f"[{timestamp}] [{level}] {message}\n")
        self.log_textbox.see("end")
        self.log_textbox.configure(state="disabled")

    def finish_process(self, final_message, severity="INFO"):
        """Finish the progress modal and report the actual scan outcome."""
        self._modal_finished = True
        self.title(t("log_modal_title_finished"))
        # Do not imply 100% completion for cancelled or invalid runs.
        self.master.reset_ui()
        self.close_button.configure(state="normal")
        self.protocol("WM_DELETE_WINDOW", self.destroy)
        if severity == "ERROR":
            messagebox.showerror(t("log_modal_finished_title"), final_message, parent=self)
        elif severity == "WARNING":
            messagebox.showwarning(t("log_modal_finished_title"), final_message, parent=self)
        else:
            messagebox.showinfo(t("log_modal_finished_title"), final_message, parent=self)
        self.grab_set()

    def on_closing_attempt(self):
        """Impede que o usuário feche o modal enquanto o processo está rodando."""
        messagebox.showwarning(t("log_modal_close_warning_title"), t("log_modal_close_warning_body"), parent=self)
        self.grab_set()

class SSHCredentialsFrame(ctk.CTkFrame if ctk is not None else object):
    """Componente reutilizável para entrada de credenciais SSH, focado em clareza e usabilidade."""
    def __init__(self, master, config: dict):
        super().__init__(master, fg_color="transparent")
        self.config = config
        self.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(self, text=t("ssh_credentials_title"), font=THEME["font_h2"]).grid(row=0, column=0, columnspan=3, sticky="w", pady=(THEME["padding_sm"], 5))

        ctk.CTkLabel(self, text=t("ssh_user_label"), font=THEME["font_body"]).grid(row=1, column=0, sticky="w")
        self.ssh_user_entry = ctk.CTkEntry(self, font=THEME["font_body"])
        self.ssh_user_entry.insert(0, self.config.get("last_ssh_user", ""))
        self.ssh_user_entry.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(0, THEME["padding_sm"]))
        Tooltip(self.ssh_user_entry, t("ssh_user_tooltip"))

        self.ssh_pass_entry = ctk.CTkEntry(self, show="*", placeholder_text=t("ssh_pass_placeholder"), font=THEME["font_body"])
        self.ssh_pass_entry.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(0, THEME["padding_sm"]))
        Tooltip(self.ssh_pass_entry, t("ssh_pass_tooltip"))

        key_frame = ctk.CTkFrame(self, fg_color="transparent")
        key_frame.grid(row=4, column=0, columnspan=3, sticky="ew")
        key_frame.grid_columnconfigure(0, weight=1)

        self.ssh_key_entry = ctk.CTkEntry(key_frame, placeholder_text=t("ssh_key_placeholder"), font=THEME["font_body"])
        self.ssh_key_entry.insert(0, self.config.get("last_ssh_key_path", ""))
        self.ssh_key_entry.grid(row=0, column=0, sticky="ew")
        Tooltip(self.ssh_key_entry, t("ssh_key_tooltip"))

        self.browse_key_button = ctk.CTkButton(key_frame, text=t("browse_button"), width=120, command=self.browse_key_file)
        self.browse_key_button.grid(row=0, column=1, padx=(10, 0))

    def browse_key_file(self):
        """Abre uma janela para o usuário selecionar o arquivo da chave SSH privada."""
        filepath = filedialog.askopenfilename(title=t("select_ssh_key_title"))
        if filepath:
            self.ssh_key_entry.delete(0, "end")
            self.ssh_key_entry.insert(0, filepath)

    def get_credentials(self) -> Dict[str, str]:
        """
        Coleta e retorna as credenciais SSH preenchidas nos campos.

        Returns:
            Um dicionário com as credenciais 'user', 'pass' e 'key_path'.
        """
        return {
            "user": self.ssh_user_entry.get().strip(),
            "pass": self.ssh_pass_entry.get(),
            "key_path": self.ssh_key_entry.get().strip()
        }

class App(ctk.CTk if ctk is not None else object):
    """
    Classe principal da aplicação que gerencia a janela, widgets e eventos.

    Orquestra a experiência do usuário, desde a tela de boas-vindas até o início
    do processo de inventário e o feedback visual através do modal de logs.
    """
    def __init__(self):
        super().__init__()
        configure_typography(self)
        self.title(APP_BRAND)
        if platform.system() == "Linux":
            try:
                self.tk.call("wm", "class", self._w, APP_NAME)
            except tk.TclError:
                pass
        self.geometry("1280x820")
        self.minsize(800, 540)
        self.resizable(True, True)
        self.configure(fg_color=THEME["canvas"])

        icon_path = resource_path(ICON_FILE)
        try:
            if platform.system() == "Windows" and os.path.exists(icon_path):
                self.iconbitmap(icon_path)
            else:
                icon_png = resource_path("brand/invent-ssh-mark-64.png")
                if os.path.exists(icon_png):
                    self._app_icon = tk.PhotoImage(file=icon_png)
                    self.iconphoto(True, self._app_icon)
        except tk.TclError:
            logging.debug("Native window icon could not be applied on this platform.", exc_info=True)

        self.is_running = False
        self.current_engine: Optional[InventoryEngine] = None
        self.log_queue = queue.Queue()
        self.config = self.load_config()
        saved_language = self.config.get("language")
        set_language(saved_language if saved_language in LANGUAGE_DISPLAY_NAMES else get_language())

        self.create_widgets()
        self.center_window()
        self.after(100, self.on_tab_change) # Garante que a UI seja atualizada na primeira exibição
        self.after(200, self.show_welcome_modal_if_needed)

    def center_window(self):
        """Open in the largest available desktop area while keeping normal resize controls."""
        self.update_idletasks()
        self.resizable(True, True)
        try:
            self.state("zoomed")
            self.update_idletasks()
            if str(self.state()).casefold() == "zoomed":
                return
        except tk.TclError:
            pass
        try:
            self.attributes("-zoomed", True)
            self.update_idletasks()
            if str(self.state()).casefold() == "zoomed" or self.winfo_width() >= int(self.winfo_screenwidth() * 0.9):
                return
        except tk.TclError:
            pass
        width = max(800, self.winfo_screenwidth() - 24)
        height = max(540, self.winfo_screenheight() - 72)
        self.geometry(f"{width}x{height}+8+8")

    def show_welcome_modal_if_needed(self):
        """Exibe o modal de boas-vindas se o usuário não o desativou."""
        if self.config.get("show_welcome_modal", True):
            welcome = WelcomeModal(self, self.config)
            self.wait_window(welcome) # Pausa a execução até o modal ser fechado

    def create_widgets(self):
        """Create a full-viewport shell with persistent navigation and responsive content."""
        self.grid_columnconfigure(0, weight=0, minsize=208)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(1, weight=1)

        header_frame = ctk.CTkFrame(
            self, corner_radius=0, fg_color=THEME["pine"], border_width=0,
        )
        header_frame.grid(row=0, column=0, columnspan=2, sticky="ew")
        header_frame.grid_columnconfigure(2, weight=1)

        try:
            self._brand_mark = load_brand_image("invent-ssh-mark-64.png", (56, 56))
            ctk.CTkLabel(header_frame, image=self._brand_mark, text="").grid(
                row=0, column=0, rowspan=2, padx=(22, 12), pady=10, sticky="w"
            )
        except (tk.TclError, OSError, RuntimeError):
            self._brand_mark = None
            ctk.CTkLabel(
                header_frame, text="●", width=48, height=48, corner_radius=14,
                fg_color=THEME["pine"], text_color=THEME["clay"],
                font=THEME["font_h1"],
            ).grid(row=0, column=0, rowspan=2, padx=(22, 12), pady=10, sticky="w")

        brand_title = ctk.CTkFrame(header_frame, fg_color="transparent")
        brand_title.grid(row=0, column=1, rowspan=2, sticky="w", pady=10)
        brand_name_line = ctk.CTkFrame(brand_title, fg_color="transparent")
        brand_name_line.pack(anchor="w", pady=(4, 0))
        ctk.CTkLabel(
            brand_name_line, text="Invent-", font=(THEME["font_family"], 22, "bold"),
            text_color=THEME["ink"],
        ).pack(side="left")
        ctk.CTkLabel(
            brand_name_line, text="SSH", font=(THEME["font_family"], 22, "bold"),
            text_color=THEME["clay"],
        ).pack(side="left")
        ctk.CTkLabel(
            brand_title, text=t("app_subtitle"), font=THEME["font_small"],
            text_color=THEME["muted"], anchor="w",
        ).pack(anchor="w", pady=(1, 0))

        ctk.CTkLabel(
            header_frame, text="LINUX   /   SSH   /   HARDWARE",
            font=THEME["font_small"], text_color=THEME["sand"],
            fg_color=THEME["paper"], corner_radius=9,
        ).grid(row=0, column=3, rowspan=2, padx=(12, 24), ipadx=12, ipady=8, sticky="e")

        self.tab_labels = {
            "home": t("tab_home"),
            "spreadsheet": t("tab_spreadsheet"),
            "settings": t("tab_settings"),
            "about": t("tab_about"),
        }
        self.tab_key_by_label = {label: key for key, label in self.tab_labels.items()}

        self.sidebar = ctk.CTkFrame(
            self, corner_radius=0, fg_color=THEME["pine"],
            border_width=1, border_color=THEME["line"], width=208,
        )
        self.sidebar.grid(row=1, column=0, sticky="nsew")
        self.sidebar.grid_propagate(False)
        self.sidebar.grid_columnconfigure(0, weight=1)
        self.sidebar.grid_rowconfigure(5, weight=1)

        ctk.CTkLabel(
            self.sidebar, text=t("home_eyebrow"), font=(THEME["font_family"], 10, "bold"),
            text_color=THEME["muted"], anchor="w",
        ).grid(row=0, column=0, sticky="ew", padx=20, pady=(24, 12))

        self.navigation_buttons = {}
        navigation = [
            ("home", "⌂", self.tab_labels["home"]),
            ("spreadsheet", "＋", self.tab_labels["spreadsheet"]),
            ("settings", "⚙", self.tab_labels["settings"]),
            ("about", "ⓘ", self.tab_labels["about"]),
        ]
        for row, (key, symbol, label) in enumerate(navigation, start=1):
            button = ctk.CTkButton(
                self.sidebar,
                text=f"{symbol}    {label}",
                anchor="w", height=45, corner_radius=10,
                font=THEME["font_body"],
                fg_color=THEME["clay"] if key == "home" else "transparent",
                hover_color=THEME["clay_hover"] if key == "home" else THEME["forest"],
                text_color=THEME["pine"] if key == "home" else THEME["ink"],
                command=lambda nav_key=key: self.select_navigation_tab(nav_key),
            )
            button.grid(row=row, column=0, padx=12, pady=4, sticky="ew")
            self.navigation_buttons[key] = button

        side_footer = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        side_footer.grid(row=6, column=0, padx=18, pady=(0, 22), sticky="sw")
        try:
            self._sidebar_mark = load_brand_image("invent-ssh-mark-64.png", (64, 64))
            ctk.CTkLabel(side_footer, image=self._sidebar_mark, text="").pack(anchor="w", pady=(0, 8))
        except (tk.TclError, OSError, RuntimeError):
            self._sidebar_mark = None
        ctk.CTkLabel(
            side_footer, text=APP_BRAND, font=THEME["font_label"],
            text_color=THEME["ink"], anchor="w",
        ).pack(anchor="w")
        ctk.CTkLabel(
            side_footer, text=f"v{APP_VERSION}", font=THEME["font_small"],
            text_color=THEME["muted"], anchor="w",
        ).pack(anchor="w", pady=(2, 0))

        self.content_frame = ctk.CTkFrame(self, fg_color=THEME["canvas"], corner_radius=0, border_width=0)
        self.content_frame.grid(row=1, column=1, sticky="nsew")
        self.content_frame.grid_columnconfigure(0, weight=1)
        self.content_frame.grid_rowconfigure(0, weight=1)

        self.tab_view = ctk.CTkTabview(
            self.content_frame, command=self.on_tab_change,
            fg_color=THEME["canvas"], border_width=0, border_color=THEME["line"],
            corner_radius=0,
            segmented_button_fg_color=THEME["paper"],
            segmented_button_selected_color=THEME["clay"],
            segmented_button_selected_hover_color=THEME["clay_hover"],
            segmented_button_unselected_color=THEME["paper"],
            segmented_button_unselected_hover_color=THEME["paper_alt"],
            text_color=THEME["pine"],
            text_color_disabled=THEME["muted"],
        )
        self.tab_view.grid(row=0, column=0, padx=0, pady=0, sticky="nsew")
        for label in self.tab_labels.values():
            self.tab_view.add(label)

        # The native tab strip is replaced by our persistent sidebar. Keeping the
        # CTkTabview content container lets existing screens remain isolated.
        self.tab_view._segmented_button.grid_remove()

        self.create_inicio_tab()
        self.create_spreadsheet_tab()
        self.create_config_tab()
        self.create_sobre_tab()

        self.run_button = ctk.CTkButton(
            self.content_frame, text=t("run_button_spreadsheet"), height=48,
            corner_radius=10, font=THEME["font_label"],
            fg_color=THEME["clay"], hover_color=THEME["clay_hover"],
            text_color=THEME["pine"], command=self.start_inventory_thread,
        )
        self.run_button.grid(
            row=1, column=0, padx=20, pady=(8, 18), sticky="ew"
        )
        Tooltip(self.run_button, t("run_button_tooltip"))
        self._update_navigation_selection("home")

    def select_navigation_tab(self, key: str) -> None:
        """Navigate to one of the four existing screens and update its active state."""
        if key not in self.tab_labels:
            return
        self.tab_view.set(self.tab_labels[key])
        self._update_ui_for_tab()

    def _update_navigation_selection(self, selected_key: str) -> None:
        """Apply a high-contrast selected state consistently across navigation items."""
        for key, button in self.navigation_buttons.items():
            selected = key == selected_key
            button.configure(
                fg_color=THEME["clay"] if selected else "transparent",
                hover_color=THEME["clay_hover"] if selected else THEME["forest"],
                text_color=THEME["pine"] if selected else THEME["ink"],
                font=(THEME["font_family"], 13, "bold" if selected else "normal"),
            )

    def select_inventory_tab(self):
        """Navigate to the one screen where the user configures a collection."""
        self.tab_view.set(self.tab_labels["spreadsheet"])
        self._update_ui_for_tab()

    def get_active_tab_key(self) -> str:
        """Retorna a chave estável ('home', 'spreadsheet', ...) da aba atualmente selecionada."""
        return self.tab_key_by_label.get(self.tab_view.get())

    def on_tab_change(self, tab_name=None):
        """Chamado quando o usuário troca de aba para atualizar a UI."""
        # 'after_idle' garante que a aba tenha sido totalmente renderizada antes de fazer alterações
        self.after_idle(self._update_ui_for_tab)

    def _update_ui_for_tab(self):
        """Atualiza o texto do botão principal e sua visibilidade com base na aba selecionada."""
        try:
            selected_key = self.get_active_tab_key()
            if selected_key:
                self._update_navigation_selection(selected_key)
            if selected_key == "spreadsheet":
                self.run_button.configure(text=t("run_button_spreadsheet"))
                self.run_button.grid()
            else:
                self.run_button.grid_remove() # Oculta o botão em abas não executáveis
        except Exception:
            # Garante que o botão seja oculto em caso de erro
            self.run_button.grid_remove()

    def create_inicio_tab(self):
        """A short, welcoming product tour with one clear next action."""
        tab = self.tab_view.tab(self.tab_labels["home"])
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)
        main_frame = ctk.CTkScrollableFrame(tab, fg_color="transparent", corner_radius=0)
        main_frame.grid(row=0, column=0, padx=18, pady=18, sticky="nsew")
        main_frame.grid_columnconfigure(0, weight=1)

        hero = ctk.CTkFrame(
            main_frame, fg_color=THEME["pine"], corner_radius=20,
            border_width=0,
        )
        hero.grid(row=0, column=0, sticky="ew", pady=(0, 22))
        hero.grid_columnconfigure(0, weight=3)
        hero.grid_columnconfigure(1, weight=2)

        ctk.CTkLabel(
            hero, text=t("home_eyebrow"), font=THEME["font_small"],
            text_color=THEME["sand"], anchor="w",
        ).grid(row=0, column=0, padx=(25, 16), pady=(25, 8), sticky="w")
        ctk.CTkLabel(
            hero, text=t("home_hero_title"), font=THEME["font_h1"],
            text_color=THEME["ink"], justify="left", anchor="w",
            wraplength=480,
        ).grid(row=1, column=0, padx=(25, 20), pady=(0, 8), sticky="ew")
        ctk.CTkLabel(
            hero, text=t("home_hero_body"), font=THEME["font_body"],
            text_color="#DCE8FB", justify="left", anchor="w",
            wraplength=500,
        ).grid(row=2, column=0, padx=(25, 25), pady=(0, 20), sticky="ew")
        ctk.CTkButton(
            hero, text=t("home_primary_cta"), command=self.select_inventory_tab,
            fg_color=THEME["clay"], hover_color=THEME["clay_hover"],
            text_color=THEME["pine"], height=43, corner_radius=11,
            font=THEME["font_label"],
        ).grid(row=3, column=0, padx=(25, 25), pady=(0, 25), sticky="w")

        sequence = ctk.CTkFrame(hero, fg_color="transparent")
        sequence.grid(row=0, column=1, rowspan=4, padx=(4, 22), pady=24, sticky="nsew")
        sequence.grid_columnconfigure(1, weight=1)
        sequence_items = [
            ("01", t("home_step1_title"), "CSV  /  XLSX"),
            ("02", t("home_step2_title"), "SSH"),
            ("03", t("home_step3_title"), "XLSX  /  HTML  /  CSV"),
        ]
        for row, (number, label, meta) in enumerate(sequence_items):
            item = ctk.CTkFrame(
                sequence, fg_color="#0D2A4B", corner_radius=12,
                border_width=1, border_color="#234872",
            )
            item.grid(row=row, column=0, columnspan=2, sticky="ew", pady=5)
            ctk.CTkLabel(
                item, text=number, font=THEME["font_label"],
                text_color=THEME["sand"], width=38,
            ).grid(row=0, column=0, rowspan=2, padx=(12, 3), pady=12)
            ctk.CTkLabel(
                item, text=label, font=THEME["font_label"],
                text_color=THEME["ink"], anchor="w",
            ).grid(row=0, column=1, padx=(0, 12), pady=(10, 0), sticky="w")
            ctk.CTkLabel(
                item, text=meta, font=THEME["font_small"],
                text_color="#B1C7E6", anchor="w",
            ).grid(row=1, column=1, padx=(0, 12), pady=(1, 10), sticky="w")

        ctk.CTkLabel(
            main_frame, text=t("home_reports_heading"),
            font=THEME["font_h2"], text_color=THEME["ink"], anchor="w",
        ).grid(row=1, column=0, sticky="w", pady=(0, 12))
        report_specs = [
            ("XLSX", t("home_reports_xlsx_title"), t("home_reports_xlsx_body"), THEME["forest"]),
            ("HTML", t("home_reports_html_title"), t("home_reports_html_body"), THEME["clay"]),
            ("CSV", t("home_reports_csv_title"), t("home_reports_csv_body"), THEME["sage"]),
        ]
        reports_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        reports_frame.grid(row=2, column=0, sticky="ew")
        report_cards = []
        for column in range(3):
            reports_frame.grid_columnconfigure(column, weight=1, uniform="home_reports")
        for index, (format_name, title, description, accent) in enumerate(report_specs):
            card = ctk.CTkFrame(
                reports_frame, fg_color=THEME["paper"], corner_radius=14,
                border_width=1, border_color=THEME["line"],
            )
            card.grid(row=0, column=index, sticky="nsew", padx=(0 if index == 0 else 6, 0 if index == 2 else 6))
            report_cards.append(card)
            ctk.CTkLabel(
                card, text=format_name, font=THEME["font_label"],
                text_color=THEME["pine"] if format_name in {"HTML", "CSV"} else THEME["ink"],
                fg_color=accent, corner_radius=8, padx=9, pady=5,
            ).pack(anchor="w", padx=14, pady=(14, 10))
            ctk.CTkLabel(
                card, text=title, font=THEME["font_label"],
                text_color=THEME["ink"], anchor="w",
            ).pack(anchor="w", padx=14, pady=(0, 4))
            ctk.CTkLabel(
                card, text=description, font=THEME["font_small"],
                text_color=THEME["muted"], anchor="w", justify="left",
                wraplength=240,
            ).pack(anchor="w", padx=14, pady=(0, 14), fill="x")
        bind_responsive_grid(reports_frame, report_cards, min_column_width=250, max_columns=3)

        ctk.CTkLabel(
            main_frame, text=t("home_privacy_note"), font=THEME["font_small"],
            text_color=THEME["muted"], anchor="w",
        ).grid(row=3, column=0, sticky="w", pady=(20, 0))

    def create_spreadsheet_tab(self):
        """Configure a scan through three calm, clearly ordered sections."""
        tab = self.tab_view.tab(self.tab_labels["spreadsheet"])
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)
        main_frame = ctk.CTkScrollableFrame(tab, fg_color="transparent", corner_radius=0)
        main_frame.grid(row=0, column=0, padx=18, pady=18, sticky="nsew")
        main_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            main_frame, text=t("inventory_intro_title"),
            font=THEME["font_h1"], text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=0, sticky="w", pady=(0, 4))
        ctk.CTkLabel(
            main_frame, text=t("inventory_intro_body"),
            font=THEME["font_body"], text_color=THEME["muted"],
            anchor="w", justify="left", wraplength=920,
        ).grid(row=1, column=0, sticky="ew", pady=(0, 18))

        file_card = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], corner_radius=16,
            border_width=1, border_color=THEME["line"],
        )
        file_card.grid(row=2, column=0, sticky="ew", pady=(0, 12))
        file_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            file_card, text=t("inventory_file_section"),
            font=THEME["font_h2"], text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=0, columnspan=2, padx=18, pady=(16, 4), sticky="w")
        ctk.CTkLabel(
            file_card, text=t("inventory_file_hint"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", wraplength=900, justify="left",
        ).grid(row=1, column=0, columnspan=2, padx=18, pady=(0, 10), sticky="w")

        self.file_entry = ctk.CTkEntry(
            file_card, placeholder_text=t("spreadsheet_file_placeholder"),
            font=THEME["font_body"], height=41,
        )
        self.file_entry.insert(0, self.config.get("last_file_path", ""))
        self.file_entry.grid(row=2, column=0, padx=(18, 8), pady=(0, 18), sticky="ew")
        Tooltip(self.file_entry, t("spreadsheet_file_tooltip"))
        browse_button = ctk.CTkButton(
            file_card, text=t("browse_button"), width=114, height=41,
            fg_color=THEME["paper_alt"], hover_color=THEME["sage_soft"],
            text_color=THEME["ink"], border_width=1, border_color=THEME["line"],
            command=self.browse_file,
        )
        browse_button.grid(row=2, column=1, padx=(0, 18), pady=(0, 18), sticky="e")
        Tooltip(browse_button, t("spreadsheet_help_tooltip"))

        ssh_card = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], corner_radius=16,
            border_width=1, border_color=THEME["line"],
        )
        ssh_card.grid(row=3, column=0, sticky="ew", pady=(0, 12))
        ssh_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            ssh_card, text=t("inventory_credentials_section"),
            font=THEME["font_h2"], text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=0, padx=18, pady=(16, 4), sticky="w")
        ctk.CTkLabel(
            ssh_card, text=t("inventory_credentials_hint"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", wraplength=900, justify="left",
        ).grid(row=1, column=0, padx=18, pady=(0, 4), sticky="w")
        self.ssh_widgets_planilha = SSHCredentialsFrame(ssh_card, self.config)
        self.ssh_widgets_planilha.grid(row=2, column=0, padx=12, pady=(0, 12), sticky="ew")

        report_card = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], corner_radius=16,
            border_width=1, border_color=THEME["line"],
        )
        report_card.grid(row=4, column=0, sticky="ew", pady=(0, 10))
        report_card.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            report_card, text=t("inventory_output_section"),
            font=THEME["font_h2"], text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=0, columnspan=2, padx=18, pady=(16, 4), sticky="w")
        ctk.CTkLabel(
            report_card, text=t("inventory_output_hint"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", wraplength=900, justify="left",
        ).grid(row=1, column=0, columnspan=2, padx=18, pady=(0, 12), sticky="w")
        ctk.CTkLabel(
            report_card, text=t("output_format_label"),
            font=THEME["font_label"], text_color=THEME["ink"],
        ).grid(row=2, column=0, padx=(18, 10), pady=(0, 18), sticky="w")
        self.spreadsheet_output_format_var = ctk.StringVar(
            value=self.config.get("report_format", self.config.get("spreadsheet_format", "XLSX"))
        )
        spreadsheet_output_menu = ctk.CTkOptionMenu(
            report_card, variable=self.spreadsheet_output_format_var,
            values=["XLSX", "HTML", "CSV"], width=160, height=38,
        )
        spreadsheet_output_menu.grid(row=2, column=1, padx=(0, 18), pady=(0, 18), sticky="w")
        Tooltip(spreadsheet_output_menu, t("spreadsheet_output_format_tooltip"))

        security_frame = ctk.CTkFrame(
            main_frame, fg_color=THEME["danger_soft"], corner_radius=12, border_width=0,
        )
        security_frame.grid(row=5, column=0, sticky="ew", pady=(0, 8))
        security_frame.grid_columnconfigure(0, weight=1)
        self.trust_unknown_hosts_var = ctk.BooleanVar(
            value=self.config.get("trust_unknown_hosts", False)
        )
        trust_checkbox = ctk.CTkCheckBox(
            security_frame, text=t("trust_unknown_hosts_label"),
            variable=self.trust_unknown_hosts_var, font=THEME["font_body"],
            fg_color=THEME["danger"], hover_color="#A63B4A",
            border_color=THEME["danger"],
        )
        trust_checkbox.grid(row=0, column=0, padx=14, pady=(12, 2), sticky="w")
        ctk.CTkLabel(
            security_frame, text=t("trust_unknown_hosts_tooltip"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", justify="left", wraplength=900,
        ).grid(row=1, column=0, padx=14, pady=(0, 8), sticky="ew")
        Tooltip(trust_checkbox, t("trust_unknown_hosts_tooltip"))
        self.include_sensitive_identifiers_var = ctk.BooleanVar(
            value=self.config.get("include_sensitive_identifiers", False)
        )
        serial_checkbox = ctk.CTkCheckBox(
            security_frame, text=t("include_sensitive_identifiers_label"),
            variable=self.include_sensitive_identifiers_var,
            font=THEME["font_body"], fg_color=THEME["clay"],
            hover_color=THEME["clay_hover"], border_color=THEME["clay"],
            text_color=THEME["ink"],
        )
        serial_checkbox.grid(row=2, column=0, padx=14, pady=(0, 2), sticky="w")
        ctk.CTkLabel(
            security_frame, text=t("include_sensitive_identifiers_tooltip"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", justify="left", wraplength=900,
        ).grid(row=3, column=0, padx=14, pady=(0, 12), sticky="ew")
        Tooltip(serial_checkbox, t("include_sensitive_identifiers_tooltip"))

    def create_config_tab(self):
        """Settings for fleet collection, with controls and explanations kept together."""
        tab = self.tab_view.tab(self.tab_labels["settings"])
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)
        main_frame = ctk.CTkScrollableFrame(tab, fg_color="transparent", corner_radius=0)
        main_frame.grid(row=0, column=0, sticky="nsew", padx=18, pady=18)
        main_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            main_frame, text=t("settings_performance_title"),
            font=THEME["font_h1"], text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=0, sticky="w", pady=(0, 14))

        perf_frame = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], border_width=1,
            border_color=THEME["line"], corner_radius=16,
        )
        perf_frame.grid(row=1, column=0, sticky="ew", pady=(0, 14))
        perf_frame.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            perf_frame, text=t("settings_workers_label"),
            font=THEME["font_label"], text_color=THEME["ink"],
        ).grid(row=0, column=0, sticky="w", padx=(18, 15), pady=(20, 5))
        self.workers_slider = ctk.CTkSlider(
            perf_frame, from_=1, to=50, number_of_steps=49,
            progress_color=THEME["sage"], button_color=THEME["clay"],
            button_hover_color=THEME["clay_hover"],
        )
        self.workers_slider.set(max(1, min(50, int(self.config.get("max_workers", 15)))))
        self.workers_slider.grid(row=0, column=1, sticky="ew", padx=(0, 12), pady=(20, 5))
        self.workers_label = ctk.CTkLabel(
            perf_frame, text=f"{int(self.workers_slider.get())}",
            width=40, font=THEME["font_label"], text_color=THEME["sage"],
        )
        self.workers_label.grid(row=0, column=2, padx=(0, 18), pady=(20, 5))
        self.workers_slider.configure(command=lambda value: self.workers_label.configure(text=f"{int(value)}"))
        Tooltip(self.workers_slider, t("settings_workers_tooltip"))
        ctk.CTkLabel(
            perf_frame, text=t("settings_workers_tooltip"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", justify="left",
        ).grid(row=1, column=0, columnspan=3, padx=18, pady=(0, 18), sticky="ew")

        ctk.CTkLabel(
            perf_frame, text=t("settings_timeout_label"),
            font=THEME["font_label"], text_color=THEME["ink"],
        ).grid(row=2, column=0, sticky="w", padx=(18, 15), pady=(18, 5))
        self.timeout_slider = ctk.CTkSlider(
            perf_frame, from_=5, to=60, number_of_steps=11,
            progress_color=THEME["sage"], button_color=THEME["clay"],
            button_hover_color=THEME["clay_hover"],
        )
        self.timeout_slider.set(max(5, min(60, int(self.config.get("ssh_timeout", 30)))))
        self.timeout_slider.grid(row=2, column=1, sticky="ew", padx=(0, 12), pady=(18, 5))
        self.timeout_label = ctk.CTkLabel(
            perf_frame, text=f"{int(self.timeout_slider.get())} s",
            width=40, font=THEME["font_label"], text_color=THEME["sage"],
        )
        self.timeout_label.grid(row=2, column=2, padx=(0, 18), pady=(18, 5))
        self.timeout_slider.configure(command=lambda value: self.timeout_label.configure(text=f"{int(value)} s"))
        Tooltip(self.timeout_slider, t("settings_timeout_tooltip"))
        ctk.CTkLabel(
            perf_frame, text=t("settings_timeout_tooltip"),
            font=THEME["font_small"], text_color=THEME["muted"],
            anchor="w", justify="left",
        ).grid(row=3, column=0, columnspan=3, padx=18, pady=(0, 18), sticky="ew")

        language_frame = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], border_width=1,
            border_color=THEME["line"], corner_radius=16,
        )
        language_frame.grid(row=2, column=0, sticky="ew")
        language_frame.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            language_frame, text=t("settings_language_title"),
            font=THEME["font_h2"], text_color=THEME["ink"],
        ).grid(row=0, column=0, columnspan=2, pady=(18, 12), padx=18, sticky="w")
        ctk.CTkLabel(
            language_frame, text=t("settings_language_label"),
            font=THEME["font_body"], text_color=THEME["ink"],
        ).grid(row=1, column=0, sticky="w", padx=(18, 12), pady=(0, 18))
        current_language = self.config.get("language", get_language())
        self.language_var = ctk.StringVar(
            value=LANGUAGE_DISPLAY_NAMES.get(current_language, LANGUAGE_DISPLAY_NAMES[DEFAULT_LANGUAGE])
        )
        language_menu = ctk.CTkOptionMenu(
            language_frame, variable=self.language_var,
            values=list(LANGUAGE_DISPLAY_NAMES.values()),
            command=self.on_language_change, width=190, height=38,
        )
        language_menu.grid(row=1, column=1, sticky="w", padx=(0, 18), pady=(0, 18))
        Tooltip(language_menu, t("settings_language_tooltip"))

    def on_language_change(self, _selected_display_name: str = None):
        """Persiste a preferência de idioma e avisa que é necessário reiniciar para aplicá-la."""
        self.save_config()
        messagebox.showinfo(t("settings_language_title"), t("settings_language_tooltip"), parent=self)

    def create_sobre_tab(self):
        """Explain the product's scope, version, and support links without clutter."""
        tab = self.tab_view.tab(self.tab_labels["about"])
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)
        main_frame = ctk.CTkScrollableFrame(tab, fg_color="transparent", corner_radius=0)
        main_frame.grid(row=0, column=0, sticky="nsew", padx=18, pady=18)
        main_frame.grid_columnconfigure(0, weight=1)

        card = ctk.CTkFrame(
            main_frame, fg_color=THEME["paper"], corner_radius=18,
            border_width=1, border_color=THEME["line"],
        )
        card.grid(row=0, column=0, sticky="ew")
        card.grid_columnconfigure(1, weight=1)
        try:
            logo = load_brand_image("invent-ssh-mark-128.png", (112, 112))
            self._about_logo = logo
            ctk.CTkLabel(card, image=logo, text="").grid(
                row=0, column=0, rowspan=2, padx=22, pady=22, sticky="n"
            )
        except (tk.TclError, OSError, RuntimeError):
            pass
        ctk.CTkLabel(
            card, text=APP_BRAND, font=THEME["font_h1"],
            text_color=THEME["ink"], anchor="w",
        ).grid(row=0, column=1, sticky="sw", padx=(0, 22), pady=(22, 0))
        ctk.CTkLabel(
            card, text=t("app_subtitle"), font=THEME["font_body"],
            text_color=THEME["muted"], anchor="w",
        ).grid(row=1, column=1, sticky="nw", padx=(0, 22), pady=(3, 20))

        fields = {
            t("about_version"): APP_VERSION,
            t("about_developed_by"): "Lucas Motta",
            t("about_license"): "MIT License",
            t("about_github"): "https://github.com/lucaswotta/invent-ssh",
            t("about_linkedin"): "https://www.linkedin.com/in/lucaswotta",
        }
        for row_index, (label_text, value_text) in enumerate(fields.items(), start=2):
            ctk.CTkLabel(
                card, text=label_text, font=THEME["font_label"],
                text_color=THEME["muted"], anchor="w",
            ).grid(row=row_index, column=0, padx=(22, 16), pady=8, sticky="w")
            if "github.com" in value_text or "linkedin.com" in value_text:
                value = ctk.CTkLabel(
                    card, text=value_text, text_color=THEME["sage"],
                    font=THEME["font_body"], cursor="hand2", anchor="w",
                )
                value.bind("<Button-1>", lambda _event, url=value_text: webbrowser.open_new_tab(url))
            else:
                value = ctk.CTkLabel(
                    card, text=value_text, font=THEME["font_body"],
                    text_color=THEME["ink"], anchor="w",
                )
            value.grid(row=row_index, column=1, padx=(0, 22), pady=8, sticky="w")

        help_button = ctk.CTkButton(
            card, text=t("about_help_button"), command=self.open_help_modal,
            fg_color=THEME["paper_alt"], hover_color=THEME["sage_soft"],
            text_color=THEME["ink"], border_width=1, border_color=THEME["line"],
            height=40,
        )
        help_button.grid(row=len(fields) + 2, column=0, columnspan=2, padx=22, pady=(20, 22), sticky="ew")

    def open_help_modal(self):
        """Abre um modal com informações sobre a precisão da coleta de dados."""
        modal = BaseModal(self, t("help_modal_title"), "650x500")
        modal.grid_rowconfigure(0, weight=1)

        textbox = ctk.CTkTextbox(modal, wrap="word", font=THEME["font_body"])
        textbox.grid(row=0, column=0, padx=THEME["padding"], pady=THEME["padding"], sticky="nsew")

        textbox.insert("1.0", t("help_modal_body"))
        textbox.configure(state="disabled")

        close_button = ctk.CTkButton(modal, text=t("help_modal_close_button"), command=modal.destroy)
        close_button.grid(row=1, column=0, padx=THEME["padding"], pady=(0, THEME["padding"]), sticky="ew")

    def create_template_spreadsheet(self):
        """Create a small, styled XLSX template using the existing lightweight workbook engine."""
        try:
            filepath = filedialog.asksaveasfilename(
                title=t("save_template_title"),
                defaultextension=".xlsx",
                initialfile="invent-ssh-assets.xlsx",
                filetypes=[(t("excel_filetype_label"), "*.xlsx")],
            )
            if not filepath:
                return

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Assets"
            headers = ["IP", "HOSTNAME", "ASSET_TAG", "LOCATION", "OWNER"]
            worksheet.append(headers)
            worksheet.append(["192.0.2.10", "linux-node-01", "ASSET-001", "Test Lab", "Infrastructure"])
            worksheet.append(["192.0.2.11", "linux-node-02", "ASSET-002", "Remote Office", "Operations"])

            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="07172D")
                cell.font = Font(bold=True, color="FFC629")
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(bottom=Side(style="medium", color="FFC629"))
            worksheet.row_dimensions[1].height = 30
            for column, width in {"A": 22, "B": 28, "C": 18, "D": 24, "E": 22}.items():
                worksheet.column_dimensions[column].width = width
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:E{worksheet.max_row}"
            worksheet.sheet_view.showGridLines = False
            workbook.properties.title = "Invent-SSH asset input template"
            workbook.properties.creator = APP_BRAND
            workbook.save(filepath)
            workbook.close()
            messagebox.showinfo(t("template_created_title"), t("template_created_body", filepath=filepath))
        except Exception as exc:
            messagebox.showerror(t("template_error_title"), t("template_error_body", error=exc))

    def browse_file(self):
        """Abre uma janela para o usuário selecionar o arquivo de planilha de entrada."""
        filepath = filedialog.askopenfilename(title=t("select_spreadsheet_title"), filetypes=[(t("spreadsheet_filetype_label"), "*.xlsx *.xlsm *.csv"), (t("all_files_label"), "*.*")])
        if filepath:
            self.file_entry.delete(0, "end")
            self.file_entry.insert(0, filepath)

    def start_inventory_thread(self):
        """
        Valida as configurações, inicia o processo de inventário em uma nova thread
        e abre o modal de logs.
        """
        if self.is_running: return

        try:
            engine_config = self.gather_config_from_ui()
            self.save_config() # Salva as configurações para a próxima sessão
            self.is_running = True
            self.run_button.configure(state="disabled", text=t("run_button_running"))
            LogModal(self)
            # Guardamos a referência do engine para poder sinalizar cancelamento
            # caso a janela seja fechada com a coleta ainda em andamento.
            self.current_engine = InventoryEngine(engine_config, self.log_queue)
            threading.Thread(target=self.current_engine.run_inventory, daemon=True).start()
        except ValueError as e:
            messagebox.showerror(t("config_error_title"), str(e), parent=self)
            self.reset_ui()
        except Exception as e:
            messagebox.showerror(t("unexpected_error_title"), t("unexpected_error_body", error=e), parent=self)
            logging.critical("Erro inesperado ao iniciar", exc_info=True)
            self.reset_ui()

    def _gather_spreadsheet_config(self) -> Dict[str, Any]:
        """Collect the asset-list and SSH settings used for this inventory."""
        filepath = self.file_entry.get().strip()
        if not filepath or not os.path.exists(filepath):
            raise ValueError(t("error_invalid_spreadsheet"))
        return {
            "mode": "Planilha",
            "filepath": filepath,
            "output_format": self.spreadsheet_output_format_var.get(),
            "trust_unknown_hosts": self.trust_unknown_hosts_var.get(),
            "include_sensitive_identifiers": self.include_sensitive_identifiers_var.get(),
        }

    def _get_active_ssh_credentials(self) -> Dict[str, str]:
        """Obtém as credenciais SSH da aba atualmente selecionada."""
        return self.ssh_widgets_planilha.get_credentials()

    def _validate_ssh_credentials(self, ssh_creds: Dict[str, str]):
        """Valida se as credenciais SSH essenciais foram fornecidas."""
        if not ssh_creds.get("user"):
            raise ValueError(t("error_ssh_user_required"))
        if not ssh_creds.get("pass") and not ssh_creds.get("key_path"):
            raise ValueError(t("error_ssh_auth_required"))

    def gather_config_from_ui(self) -> Dict[str, Any]:
        """
        Orquestra a coleta de todas as configurações da UI, valida os dados
        e retorna um dicionário de configuração para o InventoryEngine.

        Returns:
            O dicionário de configuração completo.

        Raises:
            ValueError: Se alguma configuração obrigatória estiver faltando ou for inválida.
        """
        # 1. Coleta configurações gerais
        config = {
            "max_workers": int(self.workers_slider.get()),
            "ssh_timeout": int(self.timeout_slider.get())
        }

        # Input contract: a portable CSV/Excel endpoint list, independent of any database.
        if self.get_active_tab_key() != "spreadsheet":
            raise ValueError(t("error_invalid_tab"))
        config.update(self._gather_spreadsheet_config())

        # 3. Coleta e valida as credenciais SSH
        ssh_creds = self._get_active_ssh_credentials()
        self._validate_ssh_credentials(ssh_creds)
        config.update({
            "ssh_user": ssh_creds["user"],
            "ssh_pass": ssh_creds["pass"],
            "ssh_key_path": ssh_creds["key_path"],
        })

        return config

    def _language_code_from_display_name(self, display_name: str) -> str:
        """Converte o nome exibido no seletor de idioma (ex: 'English') de volta ao código ('en')."""
        for code, name in LANGUAGE_DISPLAY_NAMES.items():
            if name == display_name:
                return code
        return DEFAULT_LANGUAGE

    def reset_ui(self):
        """Restaura o estado da UI para 'não executando'."""
        self.is_running = False
        self.current_engine = None
        self.run_button.configure(state="normal", text=t("run_button_default"))
        self.on_tab_change() # Restaura o texto dinâmico do botão

    def open_file(self, filepath: str):
        """Abre um arquivo (como o relatório final) com o programa padrão do SO."""
        try:
            if platform.system() == "Windows":
                os.startfile(filepath)
            elif platform.system() == "Darwin": # macOS
                subprocess.run(["open", filepath], check=True)
            else: # Linux
                subprocess.run(["xdg-open", filepath], check=True)
        except Exception as e:
            messagebox.showwarning(t("open_file_warning_title"), t("open_file_warning_body", filename=os.path.basename(filepath), filepath=filepath, error=e), parent=self)

    def load_config(self) -> Dict[str, Any]:
        """
        Carrega as configurações do usuário do arquivo config.json.
        Retorna um dicionário vazio se o arquivo não existir ou for inválido.
        """
        if not os.path.exists(CONFIG_FILE):
            return {}
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except (json.JSONDecodeError, FileNotFoundError):
            return {}

    def save_config(self):
        """Salva as configurações atuais da UI no arquivo config.json."""
        # Usa as credenciais da aba Planilha como padrão para salvar 'last_user' e 'last_key'
        ssh_creds = self.ssh_widgets_planilha.get_credentials()
        config_to_save = {
            "last_file_path": self.file_entry.get(),
            "report_format": self.spreadsheet_output_format_var.get(),
            "trust_unknown_hosts": self.trust_unknown_hosts_var.get(),
            "include_sensitive_identifiers": self.include_sensitive_identifiers_var.get(),
            "last_ssh_user": ssh_creds["user"],
            "last_ssh_key_path": ssh_creds["key_path"],
            "max_workers": int(self.workers_slider.get()),
            "ssh_timeout": int(self.timeout_slider.get()),
            "show_welcome_modal": self.config.get("show_welcome_modal", True),
            "language": self._language_code_from_display_name(self.language_var.get())
        }
        try:
            save_user_config(config_to_save)
        except Exception as e:
            logging.error(f"Erro ao salvar configurações: {e}")

    def on_closing(self):
        """Lida com o evento de fechamento da janela principal."""
        if self.is_running:
            if messagebox.askyesno(t("exit_confirm_title"), t("exit_confirm_body")):
                # Sinaliza cancelamento ao motor antes de destruir a janela: sem isso, as
                # threads do ThreadPoolExecutor continuam processando TODOS os terminais
                # pendentes em segundo plano (não são daemon), e o processo Python só
                # encerra de fato após elas terminarem — mesmo com a janela já fechada.
                if self.current_engine:
                    self.current_engine.cancel()
                self.destroy()
        else:
            self.save_config()
            self.destroy()

# --- Funções de Tratamento de Erro e Logging ---

def sanitize_traceback(text: str) -> str:
    """
    Remove informações sensíveis (senhas) de um traceback para
    registro seguro em logs.
    """
    # Padrões de regex para encontrar e censurar senhas e chaves
    patterns = [
        r"(password\s*=\s*['\"]).+?(['\"])",
        r"(ssh_pass\s*=\s*['\"]).+?(['\"])",
        r"(pkey\s*=\s*<paramiko\.).+?(>)"
    ]
    for p in patterns:
        text = re.sub(p, r"\1[REDACTED]\2", text, flags=re.IGNORECASE)
    return text

def _cleanup_old_files(directory: str, pattern: str, max_files: int) -> None:
    """
    Remove os arquivos mais antigos de `directory` que excedam `max_files`.
    Evita que `logs/` e `reports/` cresçam indefinidamente, já que cada
    execução gera um arquivo novo e nada os limpava antes.
    """
    if not os.path.isdir(directory):
        return
    files = sorted(glob.glob(os.path.join(directory, pattern)), key=os.path.getmtime, reverse=True)
    for old_file in files[max_files:]:
        try:
            os.remove(old_file)
        except OSError as e:
            logging.warning(f"Não foi possível remover o arquivo antigo '{old_file}': {e}")

def setup_logging():
    """Configura o sistema de logging para salvar logs detalhados em arquivos."""
    log_dir = str(get_logs_dir())
    _cleanup_old_files(log_dir, f"{APP_NAME}_*.log", max_files=20)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_file = os.path.join(log_dir, f"{APP_NAME}_{timestamp}.log")
    logging.basicConfig(
        level=logging.DEBUG,
        format='%(asctime)s - %(levelname)s - [%(module)s:%(lineno)d] - %(message)s',
        handlers=[logging.FileHandler(log_file, encoding='utf-8')]
    )
    logging.info(f"Aplicação {APP_NAME} v{APP_VERSION} iniciada.")

def main():
    """Função principal que inicializa e executa a aplicação."""
    try:
        setup_logging()
        from paths import get_reports_dir
        _cleanup_old_files(str(get_reports_dir()), "hardware_inventory_*.*", max_files=30)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme(resource_path("brand/theme.json"))
        app = App()
        app.protocol("WM_DELETE_WINDOW", app.on_closing)
        app.mainloop()
    except Exception:
        # Mecanismo de "último recurso" para capturar erros fatais
        sanitized_error = sanitize_traceback(traceback.format_exc())
        log_file = str(get_logs_dir() / "crash_log.txt")
        with open(log_file, "w", encoding='utf-8') as f:
            f.write(f"{t('fatal_error_title')}\n\n{t('fatal_error_body', log_file=log_file)}\n\n"
                    f"{'='*20} REDACTED {'='*20}\n{sanitized_error}")

        logging.critical("Erro fatal não tratado na aplicação:", exc_info=True)

        # Tenta exibir uma caixa de diálogo de erro nativa
        try:
            if tk is None or messagebox is None:
                raise RuntimeError("Tcl/Tk is unavailable.")
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror(t("fatal_error_title"), t("fatal_error_body", log_file=log_file))
        except Exception as inner_e:
            print(f"Erro fatal. Não foi possível exibir a caixa de diálogo de erro: {inner_e}")


if __name__ == "__main__":
    if check_and_install_dependencies():
        if ctk is not None:
            main()
        else:
            # Caso o usuário tenha recusado a instalação, mas as libs não existiam
            if tk is None or messagebox is None:
                raise RuntimeError("Tcl/Tk is unavailable.")
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror(t("import_error_title"), t("import_error_body"))
    sys.exit(0)