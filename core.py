# -*- coding: utf-8 -*-
"""Inventory orchestration: load endpoints and collect hardware."""
from __future__ import annotations

import concurrent.futures
import csv
import logging
import math
import os
import threading
from dataclasses import asdict, dataclass, field
from datetime import datetime
from queue import Queue
from typing import Any, Dict, List, Optional

from i18n import t
from inspector import get_hardware_info


@dataclass
class Terminal:
    """A target endpoint, its user-supplied context, and collected hardware."""
    ip: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    status: str = "PENDING"
    placa_mae: str = "N/A"
    processador: str = "N/A"
    cores_threads: str = "N/A"
    ram: str = "N/A"
    disk_type: str = "N/A"
    disk_size: str = "N/A"
    distro: str = "N/A"
    kernel: str = "N/A"
    erro: str = ""
    dta_atualizacao: Optional[datetime] = None
    hardware_details: Dict[str, Any] = field(default_factory=dict)


_ENDPOINT_ALIASES = ("IP", "HOST", "HOSTNAME", "ADDRESS", "IP_ADDRESS", "ENDPOINT", "TARGET")
# Per-machine SSH access; these cells override the SSH panel and never reach records.
_USER_ALIASES = ("USER", "USERNAME", "SSH_USER", "LOGIN", "USUARIO", "USUÁRIO")
_PASSWORD_ALIASES = ("PASSWORD", "PASS", "SSH_PASSWORD", "SENHA", "CONTRASENA", "CONTRASEÑA")


def _normalize_column_name(value: Any) -> str:
    return str(value).strip().upper().replace(" ", "_").replace("-", "_")


def _cell_value(value: Any) -> Any:
    """Normalize empty spreadsheet cells without depending on a dataframe library."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


class InventoryEngine:
    """Collect hardware with bounded concurrency and a UI-friendly event queue.

    The number of in-flight SSH tasks is bounded by max_workers; importing a
    large asset list does not create one Future per machine in memory.
    """

    def __init__(self, config: Dict[str, Any], log_queue: Queue):
        self.config = config
        self.log_queue = log_queue
        self.terminals: List[Terminal] = []
        self.load_summary: Dict[str, Any] = {}
        # Address -> {"user": ..., "password": ...} from the spreadsheet. Kept apart from
        # Terminal so records, reports, snapshots and UI events cannot carry passwords.
        self.credentials: Dict[str, Dict[str, str]] = {}
        self.logger = logging.getLogger(__name__)
        self.cancel_event = threading.Event()

    def cancel(self) -> None:
        """Request cancellation; active SSH calls finish within their timeouts."""
        self.cancel_event.set()

    def log(self, level: str, message: str, value: Any = None) -> None:
        self.log_queue.put((level, message, value))
        # Per-host records and progress are UI events; they would flood the log file.
        if level in {"HOST_RESULT", "PROGRESS"}:
            log_level = logging.DEBUG
        else:
            log_level = getattr(logging, level.upper(), logging.INFO)
        if value is None:
            self.logger.log(log_level, "%s", message)
        else:
            self.logger.log(log_level, "%s | %s", message, value)

    def run_inventory(self) -> None:
        """Run one inventory cycle and always report its real outcome to the UI."""
        final_message = t("log_process_finished", online=0, failed=0)
        final_severity = "INFO"
        try:
            self.log("INFO", t("log_starting"))
            self.terminals = self._load_from_spreadsheet()
            if not self.terminals:
                final_message = t("log_no_terminals")
                final_severity = "WARNING"
                self.log("ERROR", final_message)
                return

            self.log("INFO", t("log_terminals_loaded", count=len(self.terminals)))
            results = self._execute_collection()
            if not results:
                if self.cancel_event.is_set():
                    final_message = t("log_process_cancelled_no_results")
                else:
                    final_message = t("log_no_data_collected")
                final_severity = "WARNING"
                self.log("WARNING", final_message)
                return

            online = sum(result.status == "ONLINE" for result in results)
            failed = len(results) - online
            self.log("INFO", t("log_collection_finished", count=len(results)))

            if self.cancel_event.is_set():
                final_message = t("log_process_cancelled_with_results", count=len(results))
                final_severity = "WARNING"
            else:
                final_message = t("log_process_finished", online=online, failed=failed)
                if failed:
                    final_severity = "WARNING"
        except Exception as exc:
            final_message = t("log_process_failed", error=exc)
            final_severity = "ERROR"
            self.log("ERROR", final_message)
            self.logger.critical("Unhandled inventory engine error", exc_info=True)
        finally:
            self.log("FINISH", final_message, final_severity)

    def _load_from_spreadsheet(self) -> List[Terminal]:
        """Stream a CSV/XLSX target list while preserving arbitrary asset metadata."""
        filepath = str(self.config.get("filepath", "")).strip()
        if not filepath:
            self.log("ERROR", t("log_file_not_found", filepath=filepath))
            return []

        workbook = None
        try:
            extension = os.path.splitext(filepath)[1].lower()
            if extension in {".xlsx", ".xlsm"}:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    filepath,
                    read_only=True,
                    data_only=True,
                    keep_vba=extension == ".xlsm",
                )
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)
                return self._terminals_from_rows(headers, rows)
            if extension in {".csv", ".txt"}:
                with open(filepath, "r", encoding="utf-8-sig", newline="") as handle:
                    sample = handle.read(65536)
                    handle.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\\t|")
                    except csv.Error:
                        dialect = csv.excel
                    reader = csv.reader(handle, dialect)
                    headers = next(reader, None)
                    return self._terminals_from_rows(headers, reader)
            raise ValueError(t("loader_unsupported_format"))
        except FileNotFoundError:
            self.log("ERROR", t("log_file_not_found", filepath=filepath))
        except ValueError as exc:
            self.log("ERROR", t("log_spreadsheet_format_error", error=exc))
        except Exception as exc:
            self.log(
                "ERROR",
                t("log_spreadsheet_read_error", filename=os.path.basename(filepath), error=exc),
            )
        finally:
            if workbook is not None:
                workbook.close()
        return []

    def _terminals_from_rows(self, headers, rows) -> List[Terminal]:
        """Validate normalized headers and build records incrementally from row values."""
        if not headers:
            raise ValueError(t("loader_empty_file"))

        normalized = list(headers)
        # Ignore trailing formatting-only columns in Excel and retain meaningful interior blanks.
        while normalized and _cell_value(normalized[-1]) in (None, ""):
            normalized.pop()
        if not normalized:
            raise ValueError(t("loader_no_columns"))

        columns = [
            _normalize_column_name(name) if _cell_value(name) not in (None, "") else f"UNNAMED_{index + 1}"
            for index, name in enumerate(normalized)
        ]
        if len(set(columns)) != len(columns):
            raise ValueError(t("loader_duplicate_columns"))

        endpoint_column = next((name for name in _ENDPOINT_ALIASES if name in columns), None)
        if endpoint_column is None:
            raise ValueError(t("loader_no_target_column"))
        endpoint_index = columns.index(endpoint_column)
        user_column = next((name for name in _USER_ALIASES if name in columns), None)
        password_column = next((name for name in _PASSWORD_ALIASES if name in columns), None)
        credential_indexes = {
            columns.index(name): key
            for name, key in ((user_column, "user"), (password_column, "password"))
            if name is not None
        }

        terminals: List[Terminal] = []
        credentials: Dict[str, Dict[str, str]] = {}
        seen: set[str] = set()
        rows_without_target: List[int] = []
        duplicate_rows: List[int] = []
        for row_number, source_values in enumerate(rows, start=2):
            values = list(source_values)
            # Spreadsheet rows may be shorter than their header; pad missing cells.
            while values and _cell_value(values[-1]) in (None, ""):
                values.pop()
            if not values:
                continue
            if len(values) > len(columns):
                raise ValueError(t("loader_data_beyond_header", row=row_number))
            if len(values) < len(columns):
                values.extend([None] * (len(columns) - len(values)))

            endpoint = _cell_value(values[endpoint_index])
            endpoint = "" if endpoint is None else str(endpoint).strip()
            if not endpoint or endpoint.casefold() in {"nan", "none"}:
                rows_without_target.append(row_number)
                continue
            # Results are matched by address; a repeated one would never leave PENDING.
            if endpoint in seen:
                duplicate_rows.append(row_number)
                continue
            seen.add(endpoint)

            metadata: Dict[str, Any] = {}
            access: Dict[str, str] = {}
            for index, column in enumerate(columns):
                if index == endpoint_index:
                    continue
                value = _cell_value(values[index])
                if value is None or not str(value).strip():
                    continue
                if index in credential_indexes:
                    key = credential_indexes[index]
                    # Passwords keep their spaces; user names do not.
                    access[key] = str(value) if key == "password" else str(value).strip()
                else:
                    metadata[column] = value
            terminals.append(Terminal(ip=endpoint, metadata=metadata))
            if access:
                credentials[endpoint] = access

        credential_columns = [name for name in (user_column, password_column) if name is not None]
        self.credentials = credentials
        self.load_summary = {
            "target_column": endpoint_column,
            "extra_columns": [
                name for name in columns if name != endpoint_column and name not in credential_columns
            ],
            "credential_columns": credential_columns,
            "hosts_with_user": sum(1 for access in credentials.values() if "user" in access),
            "hosts_with_password": sum(1 for access in credentials.values() if "password" in access),
            "rows_without_target": rows_without_target,
            "duplicate_rows": duplicate_rows,
        }
        return terminals

    def _execute_collection(self) -> List[Terminal]:
        """Process thousands of endpoints without unbounded Future allocation."""
        results: List[Terminal] = []
        total = len(self.terminals)
        max_workers = max(1, min(int(self.config.get("max_workers", 15)), 128))
        timeout = max(1, int(self.config.get("ssh_timeout", 30)))
        target_iter = iter(self.terminals)
        pending: Dict[concurrent.futures.Future, Terminal] = {}
        processed = 0
        consecutive_failures = 0
        online_count = 0
        exhausted = False
        circuit_tripped = False

        executor = concurrent.futures.ThreadPoolExecutor(
            max_workers=max_workers, thread_name_prefix="hardware-inventory"
        )
        try:
            while pending or not exhausted:
                if self.cancel_event.is_set():
                    self.log("WARNING", t("log_cancelled_by_user"))
                    break

                while len(pending) < max_workers and not exhausted:
                    try:
                        terminal = next(target_iter)
                    except StopIteration:
                        exhausted = True
                        break
                    future = executor.submit(self._process_single_terminal, terminal, timeout)
                    pending[future] = terminal

                if not pending:
                    break

                done, _ = concurrent.futures.wait(
                    pending,
                    timeout=0.5,
                    return_when=concurrent.futures.FIRST_COMPLETED,
                )
                if not done:
                    continue

                for future in done:
                    terminal = pending.pop(future)
                    processed += 1
                    try:
                        result = future.result()
                        if result is not None:
                            results.append(result)
                            if result.status == "ONLINE":
                                online_count += 1
                                consecutive_failures = 0
                            else:
                                consecutive_failures += 1
                            self.log("HOST_RESULT", result.ip, asdict(result))
                    except Exception as exc:
                        terminal.status = "ERROR"
                        terminal.erro = str(exc)
                        terminal.dta_atualizacao = datetime.now()
                        results.append(terminal)
                        consecutive_failures += 1
                        self.log("HOST_RESULT", terminal.ip, asdict(terminal))
                        self.log("ERROR", t("log_processing_exception", ip=terminal.ip, error=exc))

                    self.log(
                        "PROGRESS",
                        f"{processed}/{total}: {terminal.ip}",
                        processed / total * 100 if total else 100,
                    )
                    if consecutive_failures >= 10 and online_count == 0:
                        self.log("ERROR", t("log_circuit_breaker_tripped"))
                        self.log("ERROR", t("log_circuit_breaker_hint"))
                        circuit_tripped = True
                        break

                if circuit_tripped:
                    break
        finally:
            for future in pending:
                future.cancel()
            executor.shutdown(wait=True, cancel_futures=True)

        if circuit_tripped:
            self.log("WARNING", t("log_stopped_early", count=processed))
        elif self.cancel_event.is_set():
            self.log("WARNING", t("log_cancelled_after", processed=processed, total=total))
        return results

    def _process_single_terminal(self, terminal: Terminal, timeout: int) -> Terminal:
        """Collect one endpoint while preserving the user's source metadata."""
        if not terminal.ip:
            terminal.status = "ERROR"
            terminal.erro = "Missing target address"
            terminal.dta_atualizacao = datetime.now()
            return terminal

        access = self.credentials.get(terminal.ip, {})
        hw_info = get_hardware_info(
            ip=terminal.ip,
            username=access.get("user") or self.config.get("ssh_user", ""),
            password=access.get("password") or self.config.get("ssh_pass"),
            key_path=self.config.get("ssh_key_path"),
            timeout=timeout,
            port=int(self.config.get("ssh_port", 22)),
            trust_unknown_hosts=bool(self.config.get("trust_unknown_hosts", False)),
            include_sensitive_identifiers=bool(self.config.get("include_sensitive_identifiers", False)),
        )
        raw_status = hw_info.get("status")
        if raw_status in {"SUCESSO", "SUCESSO_PARCIAL"}:
            terminal.status = "ONLINE"
            self.log("INFO", t("log_collection_success", ip=terminal.ip))
        elif raw_status == "FALHA_CONEXAO":
            terminal.status = "OFFLINE"
            terminal.erro = str(hw_info.get("erro", "Connection failed"))
            self.log("WARNING", t("log_collection_failure", ip=terminal.ip, error=terminal.erro))
        else:
            terminal.status = str(raw_status or "ERROR")
            terminal.erro = str(hw_info.get("erro", ""))
            self.log(
                "WARNING",
                t("log_collection_failure", ip=terminal.ip, error=terminal.erro or terminal.status),
            )

        for key in ("placa_mae", "processador", "cores_threads", "ram", "disk_type", "disk_size", "distro", "kernel"):
            if hw_info.get(key) not in (None, ""):
                setattr(terminal, key, hw_info[key])
        details = hw_info.get("hardware_details")
        if isinstance(details, dict):
            terminal.hardware_details = details
        terminal.dta_atualizacao = datetime.now()
        return terminal
