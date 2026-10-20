"""Professional, self-contained exports for hardware inventory runs.

The reporting layer intentionally has no dependency on databases or the GUI.
"""
from __future__ import annotations

import base64
import csv
import html
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


BRAND = "Invent-SSH"
PRODUCT_LINE = "Hardware inventory"
BRAND_NAVY = "061327"
BRAND_BLUE = "12385F"
BRAND_TEAL = "FFC629"
PALE_BLUE = "102746"
PALE_GREEN = "0D4039"
PALE_RED = "3A202C"
PALE_AMBER = "3B321A"

HARDWARE_COLUMNS = [
    ("status", "Status"),
    ("placa_mae", "Motherboard"),
    ("processador", "Processor"),
    ("cores_threads", "CPU cores / threads"),
    ("ram", "Memory"),
    ("disk_type", "Storage type"),
    ("disk_size", "Storage capacity"),
    ("distro", "Linux distribution"),
    ("kernel", "Kernel"),
    ("dta_atualizacao", "Last collected"),
    ("erro", "Details / error"),
]

HARDWARE_DETAIL_COLUMNS = [
    ("hostname", "Detected hostname"),
    ("manufacturer", "Manufacturer"),
    ("model", "System model"),
    ("architecture", "Architecture"),
    ("cpu_sockets", "CPU sockets"),
    ("cpu_max_frequency", "CPU max frequency"),
    ("cpu_virtualization", "CPU virtualization"),
    ("ram_speed", "Memory speed"),
    ("ram_slots", "Memory slots"),
    ("ram_used", "RAM used"),
    ("ram_available", "RAM available"),
    ("storage_devices", "Storage devices"),
    ("gpu", "Graphics"),
    ("network_adapters", "Network adapters"),
    ("bios", "BIOS / firmware"),
    ("uptime", "System uptime"),
    ("virtualization", "Virtualization"),
    ("system_serial", "System serial (opt-in)"),
    ("coverage_percent", "Data coverage"),
    ("collection_notes", "Collection notes"),
    ("storage_classification_basis", "Storage classification basis"),
    ("collection_sources", "Collection sources"),
]

METADATA_LABELS = {
    "HOSTNAME": "Hostname",
    "ASSET_TAG": "Asset tag",
    "ASSET_ID": "Asset ID",
    "SERIAL_NUMBER": "Serial number",
    "SERIAL": "Serial number",
    "LOCATION": "Location",
    "SITE": "Site",
    "DEPARTMENT": "Department",
    "OWNER": "Owner",
    "ENVIRONMENT": "Environment",
    "ROLE": "Role",
    "NOTES": "Notes",
    "MODEL": "Asset model",
    "MANUFACTURER": "Asset manufacturer",
    "OPERATING_SYSTEM": "Expected operating system",
}
METADATA_LABELS_PT = {
    "HOSTNAME": "Nome do host",
    "ASSET_TAG": "Etiqueta do ativo",
    "ASSET_ID": "ID do ativo",
    "SERIAL_NUMBER": "Número de série",
    "SERIAL": "Número de série",
    "LOCATION": "Localização",
    "SITE": "Unidade",
    "DEPARTMENT": "Departamento",
    "OWNER": "Responsável",
    "ENVIRONMENT": "Ambiente",
    "ROLE": "Função",
    "NOTES": "Observações",
    "MODEL": "Modelo do ativo",
    "MANUFACTURER": "Fabricante do ativo",
    "OPERATING_SYSTEM": "Sistema operacional esperado",
}
COLUMN_LABELS = {
    "en": {
        "Host / IP": "Host / IP", "Status": "Status", "Motherboard": "Motherboard",
        "Processor": "Processor", "CPU cores / threads": "CPU cores / threads",
        "Memory": "Memory", "Storage type": "Storage type", "Storage capacity": "Storage capacity",
        "Linux distribution": "Linux distribution", "Kernel": "Kernel",
        "Last collected": "Last collected", "Details / error": "Details / error",
        "Detected hostname": "Detected hostname", "Manufacturer": "Manufacturer", "System model": "System model",
        "Architecture": "Architecture", "CPU sockets": "CPU sockets",
        "CPU max frequency": "CPU max frequency", "CPU virtualization": "CPU virtualization",
        "Memory speed": "Memory speed", "Memory slots": "Memory slots", "RAM used": "RAM used",
        "RAM available": "RAM available", "Storage devices": "Storage devices", "Graphics": "Graphics",
        "Network adapters": "Network adapters", "BIOS / firmware": "BIOS / firmware",
        "System uptime": "System uptime", "Virtualization": "Virtualization",
        "System serial (opt-in)": "System serial (opt-in)", "Data coverage": "Data coverage",
        "Collection notes": "Collection notes", "Storage classification basis": "Storage classification basis",
        "Collection sources": "Collection sources",
    },
    "pt-br": {
        "Host / IP": "Host / IP", "Status": "Status", "Motherboard": "Placa-mãe",
        "Processor": "Processador", "CPU cores / threads": "Núcleos / threads da CPU",
        "Memory": "Memória", "Storage type": "Tipo de armazenamento",
        "Storage capacity": "Capacidade de armazenamento", "Linux distribution": "Distribuição Linux",
        "Kernel": "Kernel", "Last collected": "Última coleta", "Details / error": "Detalhes / erro",
        "Detected hostname": "Hostname detectado", "Manufacturer": "Fabricante", "System model": "Modelo do equipamento",
        "Architecture": "Arquitetura", "CPU sockets": "Soquetes da CPU",
        "CPU max frequency": "Frequência máxima da CPU", "CPU virtualization": "Virtualização da CPU",
        "Memory speed": "Velocidade da memória", "Memory slots": "Slots de memória", "RAM used": "RAM em uso",
        "RAM available": "RAM disponível", "Storage devices": "Unidades de armazenamento", "Graphics": "GPU / gráficos",
        "Network adapters": "Adaptadores de rede", "BIOS / firmware": "BIOS / firmware",
        "System uptime": "Tempo ligado", "Virtualization": "Ambiente virtual",
        "System serial (opt-in)": "Número de série do sistema (opcional)",
        "Data coverage": "Cobertura dos dados", "Collection notes": "Notas da coleta",
        "Storage classification basis": "Critério de classificação do armazenamento",
        "Collection sources": "Fontes utilizadas na coleta",
    },
}
REPORT_TEXT = {
    "en": {
        "title": "Hardware inventory", "kicker": "INVENTORY REPORT",
        "subtitle": "A clear snapshot of your Linux fleet, ready to review and share.",
        "generated": "Generated", "devices_inventoried": "devices inventoried",
        "summary": "Inventory summary", "total": "Total devices", "online": "Online",
        "offline": "Offline", "other": "Other / errors", "search_label": "Search inventory",
        "search_placeholder": "Search host, model, location, owner…", "all_statuses": "All statuses",
        "devices_shown": "devices shown", "footer": "Generated locally by {brand}. Inventory data stays in this report file.",
        "overview_sheet": "Overview", "inventory_sheet": "Inventory", "metric": "METRIC",
        "devices": "DEVICES", "snapshot": "devices in this snapshot",
        "coverage": "Average data coverage", "coverage_suffix": "avg",
        "inventory_hint": "Use the Inventory sheet to filter, sort, and inspect individual devices.",
        "status_options": [("ONLINE", "Online"), ("OFFLINE", "Offline"), ("FALHA_AUTH", "Authentication failed"), ("ERRO_SSH", "SSH error"), ("ERRO_DESCONHECIDO", "Unknown error"), ("ERRO_CHAVE_HOST", "Host key rejected"), ("ERRO", "Error")],
    },
    "pt-br": {
        "title": "Inventário de hardware", "kicker": "RELATÓRIO DE INVENTÁRIO",
        "subtitle": "Uma visão clara do seu parque Linux, pronta para analisar e compartilhar.",
        "generated": "Gerado em", "devices_inventoried": "máquinas inventariadas",
        "summary": "Resumo do inventário", "total": "Total de máquinas", "online": "Acessíveis",
        "offline": "Inacessíveis", "other": "Outros / erros", "search_label": "Pesquisar inventário",
        "search_placeholder": "Buscar host, modelo, local, responsável…", "all_statuses": "Todos os status",
        "devices_shown": "máquinas exibidas", "footer": "Gerado localmente pelo {brand}. Os dados permanecem neste arquivo de relatório.",
        "overview_sheet": "Resumo", "inventory_sheet": "Inventário", "metric": "INDICADOR",
        "devices": "MÁQUINAS", "snapshot": "máquinas nesta coleta",
        "coverage": "Cobertura média dos dados", "coverage_suffix": "média",
        "inventory_hint": "Use a aba Inventário para filtrar, ordenar e inspecionar as máquinas.",
        "status_options": [("ONLINE", "Acessível"), ("OFFLINE", "Inacessível"), ("FALHA_AUTH", "Falha de autenticação"), ("ERRO_SSH", "Erro SSH"), ("ERRO_DESCONHECIDO", "Erro desconhecido"), ("ERRO_CHAVE_HOST", "Chave SSH rejeitada"), ("ERRO", "Erro")],
    },
}


def _value(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return value.isoformat(sep=" ", timespec="seconds")
        except TypeError:
            return value.isoformat()
    text = str(value)
    return "" if text.lower() in {"nan", "nat", "none"} else text


def _spreadsheet_safe(value: Any) -> str:
    text = _value(value)
    # Prefix formula-like text so Excel/LibreOffice will not execute it when opening CSV/XLSX.
    candidate = text.lstrip(" \t\r\n")
    if candidate and candidate[0] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def metadata_label(key: str, language: str = "en") -> str:
    normalized = str(key).strip().upper().replace(" ", "_").replace("-", "_")
    labels = METADATA_LABELS_PT if language == "pt-br" else METADATA_LABELS
    label = labels.get(normalized, str(key).strip().replace("_", " ").replace("-", " ").title())
    return _spreadsheet_safe(label)


def _localize_detail_value(key: str, value: Any, language: str) -> str:
    """Translate explanatory labels without changing raw identifiers or numeric facts."""
    text = _value(value)
    if language != "pt-br" or not text:
        return text

    if key == "system_serial" and text == "Not collected (opt-in)":
        return "Não coletado (opcional)"
    if key == "ram_slots":
        match = re.fullmatch(r"(\d+) populated", text)
        if match:
            count = int(match.group(1))
            return f"{count} módulo(s) instalado(s)" if count != 1 else "1 módulo instalado"
    if key == "virtualization":
        return {
            "N/A": "Não disponível",
            "Not detected (physical host not independently confirmed)": "Não detectado (host físico não confirmado independentemente)",
            "Not detected": "Não detectado",
        }.get(text, text)
    if key == "cpu_virtualization":
        return {"Available": "Disponível", "Not detected": "Não detectado", "N/A": "Não disponível"}.get(text, text)
    if key == "storage_classification_basis":
        return (
            "NVMe é identificado pelo nome/transporte; HDD versus mídia não rotativa usa a flag "
            "do kernel Linux. Discos virtuais e mídia flash podem afetar essa classificação."
        )
    if key == "collection_sources":
        labels = {
            "OS:": "Sistema operacional:",
            "CPU:": "CPU:",
            "Memory:": "Memória:",
            "Storage:": "Armazenamento:",
            "System/board/firmware:": "Sistema/placa/firmware:",
            "Graphics:": "Gráficos:",
            "Network:": "Rede:",
        }
        for source, translated in labels.items():
            text = text.replace(source, translated)
        return text
    if key == "collection_notes":
        if text.startswith("Core data unavailable: "):
            fields = text.removeprefix("Core data unavailable: ").split(", ")
            labels = {
                "Linux distribution": "distribuição Linux",
                "Kernel version": "versão do kernel",
                "Machine architecture": "arquitetura da máquina",
                "CPU model": "modelo da CPU",
                "CPU topology": "topologia da CPU",
                "Usable RAM total": "memória RAM total",
                "Storage inventory": "inventário de armazenamento",
                "System vendor/model": "fabricante/modelo do sistema",
            }
            return "Dados essenciais indisponíveis: " + ", ".join(labels.get(field, field) for field in fields)
        if text.startswith("Core inventory collected; optional detail unavailable: "):
            fields = text.removeprefix("Core inventory collected; optional detail unavailable: ").split(", ")
            labels = {
                "motherboard": "placa-mãe", "DIMM speed": "velocidade dos módulos de memória",
                "GPU": "GPU", "network adapter detail": "detalhes dos adaptadores de rede",
            }
            return "Inventário principal coletado; detalhes opcionais indisponíveis: " + ", ".join(
                labels.get(field, field) for field in fields
            )
        if text.startswith("Core inventory collected."):
            return "Inventário principal coletado. Valores e classificações dependem das interfaces do kernel e das ferramentas disponíveis no sistema."
    return text


def _brand_mark_data_uri() -> str:
    """Embed the mark so HTML exports remain self-contained and work offline."""
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    mark_path = base / "brand" / "logo-mark.svg"
    try:
        mark = mark_path.read_bytes()
    except OSError:
        return ""
    encoded = base64.b64encode(mark).decode("ascii")
    return f"data:image/svg+xml;base64,{encoded}"


def _prepare(
    records: Iterable[Mapping[str, Any]], language: str = "en"
) -> tuple[list[str], list[dict[str, str]]]:
    source = [dict(record) for record in records]
    column_labels = COLUMN_LABELS[language]
    hardware_columns = [(key, column_labels[label]) for key, label in HARDWARE_COLUMNS]
    detail_columns = [(key, column_labels.get(label, label)) for key, label in HARDWARE_DETAIL_COLUMNS]
    host_label = column_labels["Host / IP"]
    reserved_labels = {
        host_label,
        *(label for _, label in hardware_columns),
        *(label for _, label in detail_columns),
    }

    metadata_columns: dict[str, str] = {}
    custom_columns: list[str] = []
    for record in source:
        metadata = record.get("metadata") or {}
        for raw_key in metadata:
            key = str(raw_key)
            if key in metadata_columns:
                continue
            label = metadata_label(key, language)
            if label in reserved_labels:
                label = f"Asset field · {label}" if language == "en" else f"Campo do ativo · {label}"
            candidate = label
            suffix = 2
            while candidate in reserved_labels or candidate in custom_columns:
                candidate = f"{label} ({suffix})"
                suffix += 1
            metadata_columns[key] = candidate
            custom_columns.append(candidate)

    columns = [
        host_label,
        *custom_columns,
        *(label for _, label in hardware_columns),
        *(label for _, label in detail_columns),
    ]
    rows: list[dict[str, str]] = []
    for record in source:
        row = {column: "" for column in columns}
        row[host_label] = _value(record.get("ip"))
        metadata = record.get("metadata") or {}
        for raw_key, value in metadata.items():
            label = metadata_columns.get(str(raw_key))
            if label:
                row[label] = _value(value)
        for key, label in hardware_columns:
            row[label] = _value(record.get(key))
        hardware_details = record.get("hardware_details") or {}
        for key, label in detail_columns:
            row[label] = _localize_detail_value(key, hardware_details.get(key), language)
        rows.append(row)
    return columns, rows


def save_report(
    records: Iterable[Mapping[str, Any]],
    output_format: str,
    output_dir: str | Path = "reports",
    language: str = "en",
) -> Path:
    """Save inventory data as styled XLSX, portable CSV, or interactive HTML."""
    language = "pt-br" if str(language).lower().startswith("pt") else "en"
    columns, rows = _prepare(records, language)
    if not rows:
        raise ValueError("There are no inventory records to export.")

    output_format = output_format.strip().upper().lstrip(".")
    extensions = {"XLSX": ".xlsx", "CSV": ".csv", "HTML": ".html"}
    if output_format not in extensions:
        raise ValueError(f"Unsupported report format: {output_format}")

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = destination / f"hardware_inventory_{timestamp}{extensions[output_format]}"

    if output_format == "CSV":
        _write_csv(path, columns, rows)
    elif output_format == "HTML":
        _write_html(path, columns, rows, language)
    else:
        _write_xlsx(path, columns, rows, language)
    return path


def _write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), delimiter=";", extrasaction="ignore")
        writer.writeheader()
        safe_rows = [
            {column: _spreadsheet_safe(row.get(column, "")) for column in columns}
            for row in rows
        ]
        writer.writerows(safe_rows)


def _average_coverage(rows: Sequence[Mapping[str, str]], coverage_label: str) -> int:
    values = []
    for row in rows:
        match = re.match(r"^\s*(\d{1,3})\s*%", row.get(coverage_label, ""))
        if match:
            values.append(min(100, int(match.group(1))))
    return round(sum(values) / len(values)) if values else 0


def _status_counts(
    rows: Sequence[Mapping[str, str]], status_label: str = "Status"
) -> tuple[int, int, int]:
    online = sum(row.get(status_label, "").upper() == "ONLINE" for row in rows)
    offline = sum(row.get(status_label, "").upper() == "OFFLINE" for row in rows)
    return online, offline, len(rows) - online - offline


def _write_html(
    path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, str]], language: str = "en"
) -> None:
    copy = REPORT_TEXT[language]
    logo_uri = _brand_mark_data_uri()
    status_label = COLUMN_LABELS[language]["Status"]
    online, offline, other = _status_counts(rows, status_label)
    coverage_label = COLUMN_LABELS[language]["Data coverage"]
    average_coverage = _average_coverage(rows, coverage_label)
    generated = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
    table_head = "".join(f"<th scope=\"col\">{html.escape(column)}</th>" for column in columns)
    body_rows = []
    for row in rows:
        status = row.get(status_label, "").upper()
        status_labels = dict(copy["status_options"])
        visible_status = status_labels.get(status, status or "UNKNOWN")
        css_status = "online" if status == "ONLINE" else "offline" if status == "OFFLINE" else "other"
        cells = []
        for column in columns:
            value = row.get(column, "")
            if column == status_label:
                cells.append(
                    f'<td class="status-cell"><span class="status-badge {css_status}">'
                    f'{html.escape(visible_status)}</span></td>'
                )
            else:
                detail_labels = {
                    COLUMN_LABELS[language]["Storage devices"],
                    COLUMN_LABELS[language]["Graphics"],
                    COLUMN_LABELS[language]["Network adapters"],
                    COLUMN_LABELS[language]["Collection notes"],
                    COLUMN_LABELS[language]["Storage classification basis"],
                    COLUMN_LABELS[language]["Collection sources"],
                }
                cell_class = ' class="detail-cell"' if column in detail_labels else ""
                cells.append(f"<td{cell_class}>{html.escape(value, quote=True)}</td>")
        body_rows.append(
            f'<tr data-status="{html.escape(status, quote=True)}">{"".join(cells)}</tr>'
        )
    escaped_brand = html.escape(BRAND)
    document = f"""<!doctype html>
<html lang="{language}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="color-scheme" content="dark">
<title>{escaped_brand} — {html.escape(copy["title"])}</title>
<style>
:root{{--ink:#edf4ff;--muted:#9db2d0;--line:#183b63;--surface:#0a1c35;--canvas:#061327;--pine:#07172d;--forest:#12385f;--sage:#6ba7e8;--clay:#ffc629;--green:#10b981;--red:#f87171;--amber:#ffc629}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--canvas);color:var(--ink);font:14px/1.55 "Avenir Next","Segoe UI",sans-serif}}
header{{background:var(--pine);color:#f4f7ff;padding:30px clamp(20px,5vw,68px) 30px}}
.brand-lockup{{display:flex;align-items:center;gap:12px;margin-bottom:18px}}
.brand-lockup img{{width:42px;height:42px;border-radius:11px;flex:none}}
.brand-name{{font-size:18px;font-weight:700;letter-spacing:-.035em}}
.brand-tag{{font-size:10px;letter-spacing:.11em;text-transform:uppercase;color:#9db2d0;margin-left:8px}}
.kicker{{font-size:11px;letter-spacing:.15em;text-transform:uppercase;color:#ffc629;font-weight:700}}
h1{{font-family:Georgia,"Times New Roman",serif;font-size:clamp(26px,4vw,40px);line-height:1.12;margin:10px 0 8px;letter-spacing:-.025em;font-weight:400}}
.sub{{color:#dce8fb;margin:0}}.meta{{margin-top:18px;color:#9db2d0;font-size:12px}}
main{{width:100%;margin:0;padding:26px clamp(16px,3vw,44px) 56px}}
.metrics{{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:12px;margin-bottom:22px}}
.metric{{background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:18px 20px}}
.metric span{{color:var(--muted);font-size:12px;font-weight:650;display:block}}.metric strong{{font-size:30px;letter-spacing:-.04em;display:block;margin-top:4px}}
.toolbar{{display:flex;gap:12px;align-items:center;margin:0 0 12px;flex-wrap:wrap}}
input,select{{font:inherit;border:1px solid #24466f;background:#081a33;border-radius:9px;padding:11px 13px;color:var(--ink)}}
input{{flex:1;min-width:220px}}select{{min-width:150px}}.result-count{{color:var(--muted);font-size:12px;margin-left:auto}}
.table-wrap{{background:#081a33;border:1px solid var(--line);border-radius:12px;overflow:auto;box-shadow:none}}
table{{width:100%;border-collapse:separate;border-spacing:0;white-space:nowrap;font-size:12px}}
td.detail-cell{{white-space:pre-line;min-width:220px;max-width:420px;line-height:1.6}}
th{{position:sticky;top:0;z-index:1;background:#102746;color:#c4d5ed;text-transform:uppercase;font-size:10px;letter-spacing:.055em;text-align:left;padding:13px 15px;border-bottom:1px solid var(--line)}}
td{{padding:12px 15px;border-bottom:1px solid #183b63;max-width:340px;overflow:hidden;text-overflow:ellipsis}}
tbody tr:nth-child(even){{background:#091b33}}tbody tr:hover{{background:#122e51}}tbody tr:last-child td{{border-bottom:0}}
.status-badge{{display:inline-flex;border-radius:999px;padding:3px 8px;font-size:10px;font-weight:800;letter-spacing:.02em}}
.status-badge.online{{background:#0d4039;color:#69e0b7}}.status-badge.offline{{background:#3a202c;color:#ff9393}}.status-badge.other{{background:#3b321a;color:#ffda73}}
footer{{padding:18px 0;color:var(--muted);font-size:11px;text-align:center}}
@media(max-width:900px){{.metrics{{grid-template-columns:repeat(3,minmax(0,1fr))}}}}
@media(max-width:720px){{.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}header{{padding-top:26px}}.metric strong{{font-size:25px}}}}
@media print{{
:root{{--ink:#07172d;--muted:#42546d;--line:#cbd7e6;--surface:#ffffff;--canvas:#ffffff}}
body{{background:#ffffff;color:#07172d}}
header{{padding:20px;background:#061327!important;color:#f4f7ff!important;print-color-adjust:exact}}
main{{width:100%;padding:16px 0}}
.toolbar{{display:none}}
.metric{{background:#ffffff;border-color:#cbd7e6;box-shadow:none}}
.metric span{{color:#42546d}}
.table-wrap{{background:#ffffff;overflow:visible;border:1px solid #cbd7e6;box-shadow:none}}
table{{white-space:normal;font-size:9px}}
th{{position:static;background:#e9eff7;color:#1c3657;print-color-adjust:exact}}
td{{color:#07172d;border-bottom:1px solid #d7e0ec;overflow:visible;max-width:none}}
tbody tr:nth-child(even){{background:#f5f8fc}}
.status-badge.online{{background:#e2f7ef;color:#0c6b51}}
.status-badge.offline{{background:#fce9ec;color:#8b2637}}
.status-badge.other{{background:#fff3d5;color:#6d4d00}}
footer{{color:#42546d}}
.metrics{{gap:5px}}
}}
</style>
</head>
<body>
<header>
  <div class="brand-lockup">{f'<img src="{logo_uri}" alt="" aria-hidden="true">' if logo_uri else ""}
    <span class="brand-name">{escaped_brand}</span>
    <span class="brand-tag">{html.escape(copy["kicker"])}</span>
  </div>
  <h1>{html.escape(copy["title"])}</h1>
  <p class="sub">{html.escape(copy["subtitle"])}</p>
  <div class="meta">{html.escape(copy["generated"])} {html.escape(generated)} &nbsp;·&nbsp; {len(rows):,} {html.escape(copy["devices_inventoried"])}</div>
</header>
<main>
  <section class="metrics" aria-label="{html.escape(copy["summary"])}">
    <div class="metric"><span>{html.escape(copy["total"])}</span><strong>{len(rows):,}</strong></div>
    <div class="metric"><span>{html.escape(copy["online"])}</span><strong>{online:,}</strong></div>
    <div class="metric"><span>{html.escape(copy["offline"])}</span><strong>{offline:,}</strong></div>
    <div class="metric"><span>{html.escape(copy["other"])}</span><strong>{other:,}</strong></div>
    <div class="metric"><span>{html.escape(copy["coverage"])}</span><strong>{average_coverage}%</strong></div>
  </section>
  <div class="toolbar">
    <input id="search" type="search" aria-label="{html.escape(copy["search_label"])}" placeholder="{html.escape(copy["search_placeholder"], quote=True)}">
    <select id="status" aria-label="{html.escape(copy["all_statuses"])}"><option value="">{html.escape(copy["all_statuses"])}</option>{''.join(f'<option value="{html.escape(value, quote=True)}">{html.escape(label)}</option>' for value, label in copy["status_options"])}</select>
    <span id="result-count" class="result-count">{len(rows):,} {html.escape(copy["devices_shown"])}</span>
  </div>
  <div class="table-wrap"><table id="inventory"><thead><tr>{table_head}</tr></thead><tbody>{''.join(body_rows)}</tbody></table></div>
  <footer>{html.escape(copy["footer"].format(brand=BRAND))}</footer>
</main>
<script>
(()=>{{const search=document.getElementById('search'),status=document.getElementById('status'),rows=[...document.querySelectorAll('#inventory tbody tr')],count=document.getElementById('result-count'),suffix={json.dumps(" " + copy["devices_shown"])};function filter(){{const q=search.value.trim().toLowerCase(),s=status.value;let visible=0;for(const row of rows){{const matchText=row.textContent.toLowerCase().includes(q),matchStatus=!s||row.dataset.status===s,show=matchText&&matchStatus;row.hidden=!show;if(show)visible++}}count.textContent=visible.toLocaleString()+suffix}}search.addEventListener('input',filter);status.addEventListener('change',filter)}})();
</script>
</body></html>"""
    path.write_text(document, encoding="utf-8")


def _write_xlsx(
    path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, str]], language: str = "en"
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        try:
            from openpyxl.drawing.image import Image as XLImage
        except ImportError:
            XLImage = None
    except ImportError as exc:
        raise RuntimeError("Excel export needs openpyxl. Install the project requirements first.") from exc

    copy = REPORT_TEXT[language]
    online, offline, other = _status_counts(rows, COLUMN_LABELS[language]["Status"])
    coverage_label = COLUMN_LABELS[language]["Data coverage"]
    average_coverage = _average_coverage(rows, coverage_label)
    generated = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
    wb = Workbook()
    overview = wb.active
    overview.title = copy["overview_sheet"]
    overview.sheet_view.showGridLines = False
    overview.column_dimensions["A"].width = 10
    overview.column_dimensions["B"].width = 32
    overview.column_dimensions["C"].width = 18
    overview.merge_cells("B1:F1")
    overview["B1"] = f"{BRAND}  /  {copy['title'].upper()}"
    overview["B1"].font = Font(name="Georgia", size=20, bold=False, color="FFFEF9")
    overview["B1"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    overview["B1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 42
    overview.merge_cells("B2:F2")
    overview["B2"] = f"{copy['generated']} {generated} · {len(rows):,} {copy['snapshot']}"
    overview["B2"].font = Font(size=10, color="9DB2D0", italic=True)
    overview.row_dimensions[2].height = 25
    if XLImage is not None:
        logo_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "brand" / "invent-ssh-mark-128.png"
        try:
            logo = XLImage(str(logo_path))
            logo.width = 36
            logo.height = 36
            overview.add_image(logo, "A1")
        except (OSError, ValueError):
            pass
    overview["B4"], overview["C4"] = copy["metric"], copy["devices"]
    for cell in (overview["B4"], overview["C4"]):
        cell.fill = PatternFill("solid", fgColor=BRAND_BLUE)
        cell.font = Font(bold=True, color="FFC629")
    summary = [
        (copy["total"], len(rows)),
        (copy["online"], online),
        (copy["offline"], offline),
        (copy["other"], other),
        (copy["coverage"], f"{average_coverage}% {copy['coverage_suffix']}"),
    ]
    for row_index, (label, count) in enumerate(summary, start=5):
        overview.cell(row_index, 2, label)
        overview.cell(row_index, 3, count)
        overview.cell(row_index, 2).font = Font(bold=True, color=BRAND_NAVY)
        overview.cell(row_index, 3).font = Font(bold=True, size=14, color=BRAND_TEAL if label == copy["online"] else BRAND_NAVY)
        overview.cell(row_index, 2).fill = PatternFill("solid", fgColor="0A1C35")
        overview.cell(row_index, 3).fill = PatternFill("solid", fgColor="0A1C35")
    overview.column_dimensions["D"].width = 18
    overview.merge_cells("B12:F12")
    overview["B12"] = copy["inventory_hint"]
    overview["B12"].font = Font(size=10, color="9DB2D0", italic=True)
    for row_cells in overview.iter_rows(min_row=1, max_row=12, min_col=1, max_col=6):
        for cell in row_cells:
            cell.fill = PatternFill("solid", fgColor=BRAND_NAVY)
    overview["A1"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    overview["A2"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    for row_index in range(5, 10):
        row_fill = "0A1C35" if row_index % 2 else "0E2440"
        label_cell = overview.cell(row_index, 2)
        value_cell = overview.cell(row_index, 3)
        label_cell.fill = PatternFill("solid", fgColor=row_fill)
        value_cell.fill = PatternFill("solid", fgColor=row_fill)
        label_cell.font = Font(bold=True, color="EDF4FF")
        value_cell.font = Font(bold=True, size=14, color="FFC629")
    for cell in (overview["B4"], overview["C4"]):
        cell.fill = PatternFill("solid", fgColor=BRAND_BLUE)
        cell.font = Font(bold=True, color="FFC629")
    overview.freeze_panes = "A4"

    sheet = wb.create_sheet(copy["inventory_sheet"])
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(columns))
    sheet.merge_cells(f"B1:{last_column}1")
    sheet["A1"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    sheet["B1"] = f"{BRAND}  /  {copy['title'].upper()}"
    sheet["B1"].font = Font(name="Georgia", size=18, bold=False, color="FFFEF9")
    sheet["B1"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    sheet["B1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 42
    if XLImage is not None:
        logo_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "brand" / "invent-ssh-mark-128.png"
        try:
            logo = XLImage(str(logo_path))
            logo.width = 36
            logo.height = 36
            sheet.add_image(logo, "A1")
        except (OSError, ValueError):
            pass
    sheet.merge_cells(f"B2:{last_column}2")
    sheet["A2"].fill = PatternFill("solid", fgColor=BRAND_NAVY)
    sheet["B2"] = f"{copy['generated']} {generated} · {len(rows):,} {copy['devices_inventoried']}"
    sheet["A2"].font = Font(size=10, color="9DB2D0", italic=True)
    for column_index, label in enumerate(columns, start=1):
        cell = sheet.cell(4, column_index, label)
        cell.font = Font(bold=True, color="FFC629", size=10)
        cell.fill = PatternFill("solid", fgColor=BRAND_BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[4].height = 30
    for row_index, row in enumerate(rows, start=5):
        for column_index, label in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column_index, _spreadsheet_safe(row.get(label, "")))
            cell.font = Font(color="EDF4FF")
            cell.fill = PatternFill("solid", fgColor="0A1C35" if row_index % 2 else "0E2440")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=label in {
                    COLUMN_LABELS[language]["Storage devices"],
                    COLUMN_LABELS[language]["Graphics"],
                    COLUMN_LABELS[language]["Network adapters"],
                    COLUMN_LABELS[language]["Collection notes"],
                    COLUMN_LABELS[language]["Storage classification basis"],
                    COLUMN_LABELS[language]["Collection sources"],
                },
            )
            if label == COLUMN_LABELS[language]["Status"]:
                status = str(row.get(label, "")).upper()
                color = PALE_GREEN if status == "ONLINE" else PALE_RED if status == "OFFLINE" else PALE_AMBER
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True, color="69E0B7" if status == "ONLINE" else "FF9393" if status == "OFFLINE" else "FFDA73")
            if label == COLUMN_LABELS[language]["Details / error"] and row.get(label):
                cell.font = Font(color="F87171")
    for column_index, label in enumerate(columns, start=1):
        values = [label] + [row.get(label, "") for row in rows[:250]]
        maximum = max((len(str(value)) for value in values), default=10)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(12, min(maximum + 2, 46))
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{max(4, 4 + len(rows))}"
    # Keep the custom navy/yellow design; built-in Excel table themes override it.
    # Auto-filter remains fully functional and avoids an unrelated default table theme.
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = "1:4"
    wb.properties.title = f"{BRAND} {copy['title']}"
    wb.properties.subject = copy["subtitle"]
    wb.properties.creator = BRAND
    wb.save(path)
